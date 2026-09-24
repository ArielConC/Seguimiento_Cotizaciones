from __future__ import annotations

import os
import re
import sqlite3
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

import po_invoices
import distributors
from parser import QuoteData, priority_for_total


APP_DIR = Path(__file__).resolve().parent


def _resolve_data_dir() -> Path:
    configured = os.environ.get("NT_QUOTE_DATA_DIR", "").strip()
    railway_volume = os.environ.get("RAILWAY_VOLUME_MOUNT_PATH", "").strip()
    if configured:
        return Path(configured)
    if railway_volume:
        return Path(railway_volume)
    if os.environ.get("RAILWAY_ENVIRONMENT_ID"):
        raise RuntimeError(
            "Persistent storage is required on Railway. Attach a volume at /data "
            "or set NT_QUOTE_DATA_DIR before starting the application."
        )
    return APP_DIR / "data"


DATA_DIR = _resolve_data_dir()
DB_PATH = DATA_DIR / "cotizaciones.db"
IMPORT_DIR = DATA_DIR / "imports"
DEFAULT_QUOTATION_ROOT = r"C:\Users\fcoar\Dropbox\Quotation"
PARSER_VERSION = 3
LOSS_REASONS = ("Lead Time", "Over Budget", "Window Shopping", "Project Cancelled", "Mismatch", "Stale Quote")
LOCAL_TIME_ZONE = ZoneInfo(os.environ.get("NT_QUOTE_TIME_ZONE", "America/Mexico_City"))
REVIEW_EVENT_TYPES = (
    "review_saved", "reviewed", "comment_added", "comment_updated",
    "status_changed", "safe_changed", "po_date_updated", "po_total_updated",
)


class ClosingConnection(sqlite3.Connection):
    def __exit__(self, exc_type, exc_value, traceback) -> bool:
        try:
            return super().__exit__(exc_type, exc_value, traceback)
        finally:
            self.close()


def local_now() -> datetime:
    return datetime.now(LOCAL_TIME_ZONE)


def today_local() -> date:
    return local_now().date()


def calendar_periods(reference: date | None = None) -> dict[str, str]:
    """Return business-calendar boundaries using the configured Mexico time zone."""
    current = reference or today_local()
    fiscal_year = current.year - 1 if current.month < 4 else current.year
    return {
        "today": current.isoformat(),
        "month_start": current.replace(day=1).isoformat(),
        "fiscal_start": date(fiscal_year, 4, 1).isoformat(),
        "time_zone": str(LOCAL_TIME_ZONE),
    }


def now_iso() -> str:
    return local_now().isoformat(timespec="seconds")


def validate_period(start: str = "", end: str = "") -> tuple[str, str]:
    start = str(start or "").strip()
    end = str(end or "").strip()
    if bool(start) != bool(end):
        raise ValueError("Start date and end date are both required")
    if not start:
        return "", ""
    start_date = date.fromisoformat(start)
    end_date = date.fromisoformat(end)
    if end_date < start_date:
        raise ValueError("End date cannot be before start date")
    return start_date.isoformat(), end_date.isoformat()


def _days_since(value: str | None) -> int:
    if not value:
        return 0
    try:
        parsed = datetime.fromisoformat(value).date() if "T" in value else date.fromisoformat(value[:10])
    except (TypeError, ValueError):
        return 0
    return max(0, (today_local() - parsed).days)


def decorate_quote(row: dict[str, Any]) -> dict[str, Any]:
    row["quote_age_days"] = _days_since(str(row.get("quote_date") or ""))
    resolved_activity = row.get("status_changed_at") if row.get("status") in {"po", "lost"} else None
    activity = resolved_activity or row.get("last_reviewed_at") or row.get("discovered_at") or row.get("updated_at")
    row["last_activity_at"] = activity
    row["activity_is_initial"] = int(row.get("status") == "pending" and not bool(row.get("last_reviewed_at")))
    row["days_since_activity"] = _days_since(str(activity or ""))
    row["overdue"] = int(row.get("status") == "pending" and not row.get("is_safe") and row["days_since_activity"] >= 15)
    return row


def connect() -> sqlite3.Connection:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(DB_PATH, timeout=30, factory=ClosingConnection)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    connection.execute("PRAGMA journal_mode = WAL")
    connection.execute("PRAGMA busy_timeout = 30000")
    return connection


def _ensure_columns(db: sqlite3.Connection, table: str, definitions: dict[str, str]) -> None:
    present = {row["name"] for row in db.execute(f"PRAGMA table_info({table})").fetchall()}
    for name, definition in definitions.items():
        if name not in present:
            db.execute(f"ALTER TABLE {table} ADD COLUMN {name} {definition}")


def backfill_review_metadata(
    db: sqlite3.Connection, quote_table: str, event_table: str, comment_table: str
) -> None:
    """Recover review timestamps created before last_reviewed_at existed."""
    allowed = {
        ("quotes", "quote_events", "quote_comments"),
        ("special_quotes", "special_quote_events", "special_quote_comments"),
    }
    if (quote_table, event_table, comment_table) not in allowed:
        raise ValueError("Unsupported review metadata tables")
    marks = ",".join("?" for _ in REVIEW_EVENT_TYPES)
    review_time = f"""COALESCE(
        (SELECT MAX(reviewed_at) FROM (
            SELECT e.created_at AS reviewed_at FROM {event_table} e
            WHERE e.quote_id=q.id AND e.event_type IN ({marks})
            UNION ALL
            SELECT c.created_at AS reviewed_at FROM {comment_table} c WHERE c.quote_id=q.id
        )),
        CASE WHEN q.status<>'pending' OR q.is_safe=1 OR trim(COALESCE(q.follow_up_type,''))<>''
                  OR trim(COALESCE(q.comment,''))<>''
             THEN COALESCE(NULLIF(q.updated_at,''),q.status_changed_at) END
    )"""
    db.execute(
        f"""UPDATE {quote_table} AS q SET last_reviewed_at={review_time}
        WHERE q.last_reviewed_at IS NULL AND {review_time} IS NOT NULL""",
        (*REVIEW_EVENT_TYPES, *REVIEW_EVENT_TYPES),
    )
    db.execute(
        f"""UPDATE {quote_table} AS q SET last_reviewed_by=(
            SELECT e.user_id FROM {event_table} e
            WHERE e.quote_id=q.id AND e.event_type IN ({marks}) AND e.user_id IS NOT NULL
            ORDER BY e.created_at DESC,e.id DESC LIMIT 1
        ) WHERE q.last_reviewed_at IS NOT NULL AND q.last_reviewed_by IS NULL""",
        REVIEW_EVENT_TYPES,
    )


def initialize() -> None:
    with connect() as db:
        db.executescript(
            """
            CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT NOT NULL);

            CREATE TABLE IF NOT EXISTS distributor_catalog (
                code TEXT PRIMARY KEY,
                company_name TEXT NOT NULL,
                normalized_name TEXT NOT NULL UNIQUE
            );

            CREATE TABLE IF NOT EXISTS distributor_unmatched (
                normalized_name TEXT PRIMARY KEY,
                company_name TEXT NOT NULL,
                first_seen_at TEXT NOT NULL,
                last_seen_at TEXT NOT NULL,
                occurrences INTEGER NOT NULL DEFAULT 1
            );

            CREATE TABLE IF NOT EXISTS quotes (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                folio TEXT NOT NULL UNIQUE,
                quote_date TEXT NOT NULL,
                receptor TEXT NOT NULL,
                distributor_agent TEXT NOT NULL DEFAULT '',
                nt_agent TEXT NOT NULL DEFAULT '',
                customer_order TEXT NOT NULL DEFAULT '',
                total_usd REAL NOT NULL,
                currency TEXT NOT NULL DEFAULT 'USD',
                priority TEXT NOT NULL CHECK (priority IN ('S', 'A', 'B', 'C')),
                status TEXT NOT NULL CHECK (status IN ('pending', 'po', 'lost')),
                follow_up_type TEXT NOT NULL DEFAULT '',
                is_safe INTEGER NOT NULL DEFAULT 0,
                po_total_usd REAL,
                comment TEXT NOT NULL DEFAULT '',
                loss_reason TEXT NOT NULL DEFAULT '',
                source_path TEXT NOT NULL,
                source_mtime REAL NOT NULL,
                discovered_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                status_changed_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS quote_events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                quote_id INTEGER NOT NULL REFERENCES quotes(id) ON DELETE CASCADE,
                event_type TEXT NOT NULL,
                from_status TEXT,
                to_status TEXT,
                follow_up_type TEXT NOT NULL DEFAULT '',
                note TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS quote_comments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                quote_id INTEGER NOT NULL REFERENCES quotes(id) ON DELETE CASCADE,
                body TEXT NOT NULL,
                user_id INTEGER,
                user_name TEXT NOT NULL DEFAULT 'Historical data',
                created_at TEXT NOT NULL,
                is_legacy INTEGER NOT NULL DEFAULT 0
            );

            CREATE TABLE IF NOT EXISTS quote_invoices (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                quote_id INTEGER NOT NULL REFERENCES quotes(id) ON DELETE CASCADE,
                invoice_date TEXT NOT NULL,
                invoice_series TEXT NOT NULL,
                invoice_number TEXT NOT NULL,
                amount REAL NOT NULL CHECK(amount>0),
                currency TEXT NOT NULL DEFAULT 'USD',
                created_by_user_id INTEGER,
                created_by_name TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL,
                active INTEGER NOT NULL DEFAULT 1,
                is_legacy INTEGER NOT NULL DEFAULT 0,
                superseded_at TEXT
            );

            CREATE TABLE IF NOT EXISTS imports (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                token TEXT NOT NULL UNIQUE,
                workspace TEXT NOT NULL,
                filename TEXT NOT NULL,
                stored_path TEXT NOT NULL,
                file_hash TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'preview',
                rows_seen INTEGER NOT NULL DEFAULT 0,
                inserted INTEGER NOT NULL DEFAULT 0,
                updated INTEGER NOT NULL DEFAULT 0,
                duplicates INTEGER NOT NULL DEFAULT 0,
                archived INTEGER NOT NULL DEFAULT 0,
                errors INTEGER NOT NULL DEFAULT 0,
                error_detail TEXT NOT NULL DEFAULT '',
                preview_json TEXT NOT NULL DEFAULT '{}',
                user_id INTEGER,
                user_name TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL,
                confirmed_at TEXT
            );

            CREATE TABLE IF NOT EXISTS daily_notes (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                note_date TEXT NOT NULL,
                note TEXT NOT NULL,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS scan_runs (
                id INTEGER PRIMARY KEY AUTOINCREMENT, started_at TEXT NOT NULL, finished_at TEXT,
                files_seen INTEGER NOT NULL DEFAULT 0, inserted INTEGER NOT NULL DEFAULT 0,
                updated INTEGER NOT NULL DEFAULT 0, errors INTEGER NOT NULL DEFAULT 0,
                error_detail TEXT NOT NULL DEFAULT ''
            );
            CREATE TABLE IF NOT EXISTS processed_files (
                source_path TEXT PRIMARY KEY, source_mtime REAL NOT NULL, parser_version INTEGER NOT NULL,
                outcome TEXT NOT NULL, detail TEXT NOT NULL DEFAULT '', processed_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_quotes_status ON quotes(status);
            CREATE INDEX IF NOT EXISTS idx_quotes_agent ON quotes(nt_agent);
            CREATE INDEX IF NOT EXISTS idx_quotes_date ON quotes(quote_date);
            CREATE INDEX IF NOT EXISTS idx_events_created ON quote_events(created_at);
            CREATE INDEX IF NOT EXISTS idx_quote_invoices_active ON quote_invoices(quote_id,active);
            CREATE INDEX IF NOT EXISTS idx_imports_workspace ON imports(workspace, created_at);
            CREATE INDEX IF NOT EXISTS idx_distributor_catalog_name ON distributor_catalog(normalized_name);
            """
            
        )
        db.execute("""
            CREATE TABLE IF NOT EXISTS po_invoices (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                quote_id INTEGER NOT NULL,
                invoice_date TEXT,
                invoice_series TEXT,
                invoice_number TEXT,
                invoice_amount REAL,
                FOREIGN KEY(quote_id) REFERENCES quotes(id)
            )
        """)
        _ensure_columns(db, "quotes", {
            "follow_up_type": "TEXT NOT NULL DEFAULT ''", "is_safe": "INTEGER NOT NULL DEFAULT 0",
            "po_total_usd": "REAL", "series": "TEXT NOT NULL DEFAULT ''",
            "folio_number": "TEXT NOT NULL DEFAULT ''", "distributor_company": "TEXT NOT NULL DEFAULT ''",
            "end_user": "TEXT NOT NULL DEFAULT ''", "net_total_usd": "REAL", "po_date": "TEXT",
            "po_detected": "INTEGER NOT NULL DEFAULT 0", "source_kind": "TEXT NOT NULL DEFAULT 'pdf'",
            "is_historical": "INTEGER NOT NULL DEFAULT 0", "is_archived": "INTEGER NOT NULL DEFAULT 0",
            "import_id": "INTEGER", "created_by_user_id": "INTEGER",
            "created_by_name": "TEXT NOT NULL DEFAULT ''", "last_reviewed_at": "TEXT",
            "last_reviewed_by": "INTEGER",
            "distributor_code": "TEXT NOT NULL DEFAULT ''",
            "client_response": "TEXT NOT NULL DEFAULT 'pending'", # <--- NUEVA COLUMNA
        })
        _ensure_columns(db, "quote_events", {
            "follow_up_type": "TEXT NOT NULL DEFAULT ''", "user_id": "INTEGER",
            "user_name": "TEXT NOT NULL DEFAULT 'Historical data'",
        })
        db.execute("INSERT OR IGNORE INTO settings(key,value) VALUES('quotation_root',?)", (DEFAULT_QUOTATION_ROOT,))
        db.execute("INSERT OR IGNORE INTO settings(key,value) VALUES('months_back','3')")
        db.execute("INSERT OR IGNORE INTO settings(key,value) VALUES('scan_interval_seconds','300')")
        for code, company_name in distributors.CATALOG:
            db.execute(
                """INSERT INTO distributor_catalog(code,company_name,normalized_name) VALUES(?,?,?)
                ON CONFLICT(code) DO UPDATE SET company_name=excluded.company_name,
                normalized_name=excluded.normalized_name""",
                (code, company_name, distributors.normalize_name(company_name)),
            )
        catalog = {
            row["normalized_name"]: row["code"]
            for row in db.execute("SELECT code,normalized_name FROM distributor_catalog").fetchall()
        }
        for quote in db.execute("SELECT id,distributor_company,distributor_code FROM quotes").fetchall():
            code = catalog.get(distributors.normalize_name(quote["distributor_company"]))
            if code and code != quote["distributor_code"]:
                db.execute("UPDATE quotes SET distributor_code=? WHERE id=?", (code, quote["id"]))
        db.execute(
            """
            UPDATE quotes SET
                series=CASE WHEN series='' AND instr(folio,'-')>0 THEN substr(folio,1,instr(folio,'-')-1) ELSE series END,
                folio_number=CASE WHEN folio_number='' AND instr(folio,'-')>0 THEN substr(folio,instr(folio,'-')+1) ELSE folio_number END,
                end_user=CASE WHEN end_user='' THEN receptor ELSE end_user END,
                source_kind=COALESCE(NULLIF(source_kind,''),'pdf'),
                priority=CASE WHEN total_usd<=500 THEN 'C' WHEN total_usd<=1000 THEN 'B' WHEN total_usd<=5000 THEN 'A' ELSE 'S' END
            """
        )
        po_invoices.migrate_legacy(db,"quotes","quote_invoices","quote_events","USD")
        db.execute("UPDATE quotes SET po_detected=1 WHERE trim(customer_order)<>''")
        db.execute("UPDATE quotes SET status='pending' WHERE status='po' AND po_date IS NULL")
        db.execute(
            """
            INSERT INTO quote_comments(quote_id,body,user_name,created_at,is_legacy)
            SELECT q.id,q.comment,'Historical data',q.updated_at,1 FROM quotes q
            WHERE trim(q.comment)<>'' AND NOT EXISTS(
                SELECT 1 FROM quote_comments c WHERE c.quote_id=q.id AND c.is_legacy=1
            )
            """
        )
        backfill_review_metadata(db, "quotes", "quote_events", "quote_comments")


def get_setting(key: str, default: str = "") -> str:
    with connect() as db:
        row = db.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
    return row["value"] if row else default


def set_setting(key: str, value: str) -> None:
    with connect() as db:
        db.execute("INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (key, value))


def distributor_code_for(db: sqlite3.Connection, company_name: str) -> str:
    normalized = distributors.normalize_name(company_name)
    if not normalized:
        return ""
    row = db.execute("SELECT code FROM distributor_catalog WHERE normalized_name=?", (normalized,)).fetchone()
    if row:
        db.execute("DELETE FROM distributor_unmatched WHERE normalized_name=?", (normalized,))
        return str(row["code"])
    timestamp = now_iso()
    db.execute(
        """INSERT INTO distributor_unmatched(normalized_name,company_name,first_seen_at,last_seen_at,occurrences)
        VALUES(?,?,?,?,1) ON CONFLICT(normalized_name) DO UPDATE SET
        company_name=excluded.company_name,last_seen_at=excluded.last_seen_at,occurrences=distributor_unmatched.occurrences+1""",
        (normalized, str(company_name or "").strip(), timestamp, timestamp),
    )
    return ""


def rows_to_dicts(rows: list[sqlite3.Row]) -> list[dict[str, Any]]:
    return [dict(row) for row in rows]


def _actor(actor: dict[str, Any] | None) -> tuple[int | None, str]:
    return (int(actor["id"]) if actor and actor.get("id") else None, str(actor.get("display_name") if actor else "System"))


def _parse_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise ValueError(f"Unrecognized date: {text or 'blank'}")


def _number(value: Any, label: str, required: bool = True) -> float | None:
    if value in (None, "") and not required:
        return None
    try:
        result = float(str(value).replace(",", "").replace("$", "").strip())
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} must be numeric") from exc
    if result < 0:
        raise ValueError(f"{label} cannot be negative")
    return result


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_folio(series: str, number: str) -> tuple[str, str, str]:
    series = re.sub(r"\s+", "", series.upper())
    number = _cell_text(number).strip()
    number = re.sub(rf"^{re.escape(series)}[-_ ]*", "", number, flags=re.I) if series else number
    if not series or not number:
        raise ValueError("Series and folio are required")
    return series, number, f"{series}-{number}"


def parse_workbook(content: bytes, filename: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Upload an .xlsx Excel file")
    if len(content) > 25 * 1024 * 1024:
        raise ValueError("The Excel file cannot exceed 25 MB")
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError("The Excel workbook could not be read") from exc
    worksheet = workbook.active
    worksheet_rows = list(worksheet.iter_rows(min_col=1, max_col=11, values_only=True))
    header_row = None
    for row_number, row_values in enumerate(worksheet_rows[:15], start=1):
        a = _cell_text(row_values[0]).lower()
        b = _cell_text(row_values[1]).lower()
        c = _cell_text(row_values[2]).lower()
        if a in {"fecha", "date"} and b in {"serie", "series"} and c in {"folio", "quote no.", "quote no"}:
            header_row = row_number
            break
    if header_row is None:
        raise ValueError("Required headers were not found in columns A, B, and C")
    records: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for row_number, row_values in enumerate(worksheet_rows[header_row:], start=header_row + 1):
        values = list(row_values)
        if not any(value not in (None, "") for value in values):
            continue
        if values[1] in (None, "") or values[2] in (None, ""):
            continue
        try:
            series, folio_number, folio = _normalize_folio(_cell_text(values[1]), _cell_text(values[2]))
            quote_date = _parse_date(values[0])
            total = float(_number(values[4], "Total with taxes") or 0)
            net_total = _number(values[10], "Net total", required=False)
            customer_order = _cell_text(values[8])
            records.append({
                "series": series, "folio_number": folio_number, "folio": folio,
                "quote_date": quote_date.isoformat(), "distributor_company": _cell_text(values[3]),
                "end_user": _cell_text(values[5]), "receptor": _cell_text(values[5]) or _cell_text(values[3]),
                "distributor_agent": _cell_text(values[6]), "customer_order": customer_order,
                "po_detected": int(bool(customer_order)), "nt_agent": _cell_text(values[9]) or "Unassigned",
                "total_usd": total, "net_total_usd": net_total, "currency": "USD",
                "priority": priority_for_total(Decimal(str(total))), "source_row": row_number,
            })
        except Exception as exc:
            errors.append({"row": row_number, "error": str(exc)})
    if not records and not errors:
        raise ValueError("The workbook does not contain quotation rows")
    return records, errors


def preview_workbook(content: bytes, filename: str) -> dict[str, Any]:
    records, errors = parse_workbook(content, filename)
    unique: dict[str, dict[str, Any]] = {}
    duplicates = 0
    for record in records:
        if record["folio"] in unique:
            duplicates += 1
        unique[record["folio"]] = record
    with connect() as db:
        existing = {row["folio"] for row in db.execute(
            f"SELECT folio FROM quotes WHERE folio IN ({','.join('?' for _ in unique)})", tuple(unique)
        ).fetchall()} if unique else set()
    return {
        "rows_seen": len(records), "valid": len(unique), "new": len(set(unique) - existing),
        "updated": len(set(unique) & existing), "duplicates": duplicates, "archived": 0,
        "errors": len(errors), "error_detail": errors[:20], "sample": list(unique.values())[:8],
    }


def import_workbook(content: bytes, filename: str, actor: dict[str, Any] | None = None, import_id: int | None = None) -> dict[str, Any]:
    records, errors = parse_workbook(content, filename)
    unique: dict[str, dict[str, Any]] = {}
    duplicates = 0
    for record in records:
        if record["folio"] in unique:
            duplicates += 1
        unique[record["folio"]] = record
    timestamp = now_iso()
    actor_id, actor_name = _actor(actor)
    inserted = updated = 0
    with connect() as db:
        db.execute("BEGIN IMMEDIATE")
        for record in unique.values():
            distributor_code = distributor_code_for(db, record["distributor_company"])
            current = db.execute("SELECT * FROM quotes WHERE folio=?", (record["folio"],)).fetchone()
            if current:
                db.execute(
                    """
                    UPDATE quotes SET series=?,folio_number=?,quote_date=?,receptor=?,distributor_company=?,distributor_code=?,end_user=?,
                        distributor_agent=?,nt_agent=?,customer_order=?,po_detected=?,total_usd=?,net_total_usd=?,
                        currency='USD',priority=?,source_path=?,source_mtime=0,source_kind='excel',is_historical=0,
                        is_archived=0,import_id=? WHERE id=?
                    """,
                    (record["series"],record["folio_number"],record["quote_date"],record["receptor"],
                     record["distributor_company"],distributor_code,record["end_user"],record["distributor_agent"],record["nt_agent"],
                     record["customer_order"],record["po_detected"],record["total_usd"],record["net_total_usd"],
                     record["priority"],filename,import_id,current["id"]),
                )
                updated += 1
            else:
                cursor = db.execute(
                    """
                    INSERT INTO quotes(folio,series,folio_number,quote_date,receptor,distributor_company,distributor_code,end_user,
                        distributor_agent,nt_agent,customer_order,po_detected,total_usd,net_total_usd,currency,
                        priority,status,source_path,source_mtime,source_kind,import_id,created_by_user_id,created_by_name,
                        discovered_at,updated_at,status_changed_at)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,'USD',?,'pending',?,0,'excel',?,?,?,?,?,?)
                    """,
                    (record["folio"],record["series"],record["folio_number"],record["quote_date"],record["receptor"],
                     record["distributor_company"],distributor_code,record["end_user"],record["distributor_agent"],record["nt_agent"],
                     record["customer_order"],record["po_detected"],record["total_usd"],record["net_total_usd"],
                     record["priority"],filename,import_id,actor_id,actor_name,timestamp,timestamp,timestamp),
                )
                db.execute(
                    "INSERT INTO quote_events(quote_id,event_type,to_status,note,user_id,user_name,created_at) VALUES(?,'created','pending','Imported from Excel',?,?,?)",
                    (cursor.lastrowid, actor_id, actor_name, timestamp),
                )
                inserted += 1
        db.execute("UPDATE quotes SET is_historical=1 WHERE source_kind='pdf'")
    return {"rows_seen":len(records),"inserted":inserted,"updated":updated,"duplicates":duplicates,
            "archived":0,"errors":len(errors),"error_detail":errors[:20]}


def upsert_quote(data: QuoteData, source_mtime: float) -> str:
    """Legacy PDF import retained only for migration compatibility."""
    timestamp = now_iso()
    parts = data.folio.replace("_", "-").split("-", 1)
    series, number, folio = _normalize_folio(parts[0], parts[-1])
    priority = priority_for_total(Decimal(data.total_usd))
    with connect() as db:
        db.execute("BEGIN IMMEDIATE")
        current = db.execute("SELECT * FROM quotes WHERE folio=?", (folio,)).fetchone()
        if not current:
            cursor = db.execute(
                """
                INSERT INTO quotes(folio,series,folio_number,quote_date,receptor,end_user,distributor_agent,nt_agent,
                    customer_order,po_detected,total_usd,currency,priority,status,source_path,source_mtime,source_kind,
                    discovered_at,updated_at,status_changed_at)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,'USD',?,'pending',?,?,'pdf',?,?,?)
                """,
                (folio,series,number,data.quote_date.isoformat(),data.receptor,data.receptor,data.distributor_agent,
                 data.nt_agent,data.customer_order,int(bool(data.customer_order)),float(data.total_usd),priority,
                 data.source_path,source_mtime,timestamp,timestamp,timestamp),
            )
            db.execute("INSERT INTO quote_events(quote_id,event_type,to_status,note,created_at) VALUES(?,'created','pending','Historical PDF import',?)", (cursor.lastrowid,timestamp))
            return "inserted"
        if source_mtime < float(current["source_mtime"]):
            return "skipped_duplicate"
        db.execute(
            """UPDATE quotes SET quote_date=?,receptor=?,end_user=?,distributor_agent=?,nt_agent=?,customer_order=?,
            po_detected=?,total_usd=?,priority=?,source_path=?,source_mtime=? WHERE id=?""",
            (data.quote_date.isoformat(),data.receptor,data.receptor,data.distributor_agent,data.nt_agent,data.customer_order,
             int(bool(data.customer_order)),float(data.total_usd),priority,data.source_path,source_mtime,current["id"]),
        )
        return "updated"


def source_is_current(source_path: str, source_mtime: float) -> bool:
    with connect() as db:
        row = db.execute("SELECT source_mtime,parser_version FROM processed_files WHERE source_path=?", (str(Path(source_path).resolve()),)).fetchone()
    return bool(row and int(row["parser_version"]) == PARSER_VERSION and abs(float(row["source_mtime"]) - source_mtime) < .001)


def record_file_result(source_path: str, source_mtime: float, outcome: str, detail: str = "") -> None:
    with connect() as db:
        db.execute(
            """INSERT INTO processed_files(source_path,source_mtime,parser_version,outcome,detail,processed_at)
            VALUES(?,?,?,?,?,?) ON CONFLICT(source_path) DO UPDATE SET source_mtime=excluded.source_mtime,
            parser_version=excluded.parser_version,outcome=excluded.outcome,detail=excluded.detail,processed_at=excluded.processed_at""",
            (str(Path(source_path).resolve()),source_mtime,PARSER_VERSION,outcome,detail[:1000],now_iso()),
        )


def start_scan() -> int:
    with connect() as db:
        return int(db.execute("INSERT INTO scan_runs(started_at) VALUES(?)", (now_iso(),)).lastrowid)


def finish_scan(scan_id: int, stats: dict[str, Any], errors: list[str]) -> None:
    with connect() as db:
        db.execute("UPDATE scan_runs SET finished_at=?,files_seen=?,inserted=?,updated=?,errors=?,error_detail=? WHERE id=?",
                   (now_iso(),stats["files_seen"],stats["inserted"],stats["updated"],stats["errors"],"\n".join(errors[:50]),scan_id))


def _quote_where(filters: dict[str, str], alias: str = "q") -> tuple[list[str], list[Any]]:
    if filters.get("archive") == "1":
        clauses = [f"{alias}.is_archived=1"]
    else:
        clauses = [f"{alias}.is_archived=0"]
        clauses.append(f"{alias}.is_historical={'1' if filters.get('historical') == '1' else '0'}")
    values: list[Any] = []
    if filters.get("status"):
        clauses.append(f"{alias}.status=?"); values.append(filters["status"])
    if filters.get("agent"):
        clauses.append(f"{alias}.nt_agent=?"); values.append(filters["agent"])
    if filters.get("priority"):
        clauses.append(f"{alias}.priority=?"); values.append(filters["priority"])
    if filters.get("distributor"):
        needle = f"%{filters['distributor']}%"
        clauses.append(f"({alias}.distributor_code LIKE ? OR {alias}.distributor_company LIKE ?)")
        values.extend([needle, needle])
    if filters.get("end_user"):
        clauses.append(f"{alias}.end_user LIKE ?")
        values.append(f"%{filters['end_user']}%")
    if filters.get("safe") in {"0", "1"}:
        clauses.append(f"{alias}.is_safe=?"); values.append(int(filters["safe"]))
    if filters.get("po_missing") == "1":
        clauses.append(f"{alias}.po_detected=1 AND {alias}.po_date IS NULL AND {alias}.status='pending'")
    start, end = validate_period(filters.get("start", ""), filters.get("end", ""))
    if start:
        clauses.append(f"{alias}.quote_date BETWEEN ? AND ?")
        values.extend([start, end])
    if filters.get("search"):
        needle = f"%{filters['search']}%"
        clauses.append(f"({alias}.folio LIKE ? OR {alias}.distributor_code LIKE ? OR {alias}.distributor_company LIKE ? OR {alias}.end_user LIKE ? OR {alias}.receptor LIKE ? OR {alias}.distributor_agent LIKE ? OR {alias}.customer_order LIKE ?)")
        values.extend([needle] * 7)
    return clauses, values


def list_quotes(filters: dict[str, str]) -> list[dict[str, Any]]:
    clauses, values = _quote_where(filters)
    n = "CAST(q.folio_number AS INTEGER)"
    orders = {
        "priority": "CASE q.priority WHEN 'S' THEN 1 WHEN 'A' THEN 2 WHEN 'B' THEN 3 ELSE 4 END,q.total_usd DESC",
        "date_desc": f"q.quote_date DESC,{n} DESC", "date_asc": f"q.quote_date ASC,{n} ASC",
        "folio_desc": f"{n} DESC,q.quote_date DESC", "folio_asc": f"{n} ASC,q.quote_date ASC",
        "amount_desc": "q.total_usd DESC,q.quote_date DESC", "amount_asc": "q.total_usd ASC,q.quote_date ASC",
    }
    order = orders.get(filters.get("order", "priority"), orders["priority"])
    with connect() as db:
        # ---> LÍNEA MODIFICADA <---
        rows = db.execute(f"SELECT q.*, EXISTS(SELECT 1 FROM po_invoices WHERE quote_id=q.id) AS has_invoices FROM quotes q WHERE {' AND '.join(clauses)} ORDER BY {order}", values).fetchall()
    return [decorate_quote(row) for row in rows_to_dicts(rows)]


def list_managed_quotes(filters: dict[str, str] | None = None) -> list[dict[str, Any]]:
    filters = dict(filters or {})
    
    # Atrapamos el filtro de 'No gestionadas'
    unmanaged = False
    if filters.get("status") == "unmanaged":
        unmanaged = True
        filters["status"] = "pending"
        
    clauses, values = _quote_where(filters)
    clauses.append("q.last_reviewed_at IS NULL" if unmanaged else "q.last_reviewed_at IS NOT NULL")
        
    orders = {
        "priority": "CASE q.priority WHEN 'S' THEN 1 WHEN 'A' THEN 2 WHEN 'B' THEN 3 ELSE 4 END,q.total_usd DESC",
        "date_desc": "q.quote_date DESC,CAST(q.folio_number AS INTEGER) DESC",
        "date_asc": "q.quote_date ASC,CAST(q.folio_number AS INTEGER) ASC",
        "status_activity": "CASE q.status WHEN 'pending' THEN 1 WHEN 'lost' THEN 2 WHEN 'po' THEN 3 ELSE 4 END,datetime(COALESCE(q.last_reviewed_at,q.discovered_at)) DESC",
        "newest_activity": "datetime(COALESCE(q.last_reviewed_at,q.discovered_at)) DESC",
        "oldest_activity": "datetime(COALESCE(q.last_reviewed_at,q.discovered_at)) ASC",
    }
    order = orders.get(filters.get("order", "status_activity"), orders["status_activity"])
    
    with connect() as db:
        # Incluimos la validación de po_invoices para encender el gafete de Excel
        rows = db.execute(f"SELECT q.*, EXISTS(SELECT 1 FROM po_invoices WHERE quote_id=q.id) AS has_invoices FROM quotes q WHERE {' AND '.join(clauses)} ORDER BY {order}", values).fetchall()
        
    return [decorate_quote(row) for row in rows_to_dicts(rows)]

def list_management_quotes(filters: dict[str, str] | None = None) -> list[dict[str, Any]]:
    filters = dict(filters or {})
    filters.pop("historical", None)
    filters.pop("archive", None)
    
    # Atrapamos el filtro de 'No gestionadas'
    unmanaged = False
    if filters.get("status") == "unmanaged":
        unmanaged = True
        filters["status"] = "pending"
        
    clauses, values = _quote_where(filters)
    
    # Exigimos a la base de datos que la fecha de revisión esté vacía
    if unmanaged:
        clauses.append("q.last_reviewed_at IS NULL")
        
    orders = {
        "priority": "CASE q.priority WHEN 'S' THEN 1 WHEN 'A' THEN 2 WHEN 'B' THEN 3 ELSE 4 END,q.total_usd DESC",
        "date_desc": "q.quote_date DESC,CAST(q.folio_number AS INTEGER) DESC",
        "date_asc": "q.quote_date ASC,CAST(q.folio_number AS INTEGER) ASC",
        "folio_desc": "CAST(q.folio_number AS INTEGER) DESC,q.quote_date DESC",
        "folio_asc": "CAST(q.folio_number AS INTEGER) ASC,q.quote_date ASC",
        "amount_desc": "q.total_usd DESC,q.quote_date DESC",
        "amount_asc": "q.total_usd ASC,q.quote_date ASC",
        "newest_activity": "datetime(COALESCE(q.last_reviewed_at,q.discovered_at)) DESC",
        "oldest_activity": "datetime(COALESCE(q.last_reviewed_at,q.discovered_at)) ASC",
    }
    order = orders.get(filters.get("order", "newest_activity"), orders["newest_activity"])
    
    with connect() as db:
        # Incluimos la validación de po_invoices para encender el gafete de Excel
        rows = db.execute(f"SELECT q.*, EXISTS(SELECT 1 FROM po_invoices WHERE quote_id=q.id) AS has_invoices FROM quotes q WHERE {' AND '.join(clauses)} ORDER BY {order}", values).fetchall()
        
    return [decorate_quote(row) for row in rows_to_dicts(rows)]

def get_quote(quote_id: int) -> dict[str, Any]:
    with connect() as db:
        row = db.execute("SELECT * FROM quotes WHERE id=?", (quote_id,)).fetchone()
        if not row:
            raise KeyError("Quotation not found")
        comments = rows_to_dicts(db.execute("SELECT * FROM quote_comments WHERE quote_id=? ORDER BY created_at,id", (quote_id,)).fetchall())
        events = rows_to_dicts(db.execute("SELECT * FROM quote_events WHERE quote_id=? ORDER BY created_at,id", (quote_id,)).fetchall())
        invoices = po_invoices.active(db,"quote_invoices",quote_id)
    result = decorate_quote(dict(row)); result["comments"] = comments; result["events"] = events; result["invoices"] = invoices
    result["read_only"] = bool(result["is_historical"])
    return result


def update_quote(quote_id: int, status: str, comment: str, loss_reason: str, follow_up_type: str,
                 is_safe: bool, po_total_usd: float | None, po_date: str | None = None,
                 actor: dict[str, Any] | None = None, invoices: list[dict[str, Any]] | None = None,
                 client_response: str = "pending") -> dict[str, Any]:
    if status not in {"pending", "po", "lost"}:
        raise ValueError("Invalid status")
    if follow_up_type not in {"email", "call", "visit"}:
        raise ValueError("Select a follow-up method: E-mail, Call, or Visit")
    loss_reason = loss_reason.strip()
    if status == "lost" and loss_reason not in LOSS_REASONS:
        raise ValueError("Select a valid loss reason")
    if status != "lost": loss_reason = ""
    if client_response not in {"pending", "yes", "no"}:
        raise ValueError("Invalid client response")
    normalized_invoices = po_invoices.normalize(
        invoices,"USD",legacy_total=po_total_usd,legacy_date=po_date,legacy_number=f"PO-{quote_id}"
    ) if status == "po" else []
    if status == "po": po_total_usd,po_date=po_invoices.totals(normalized_invoices)
    else: po_total_usd=None; po_date=None
    comment = "" if status=="po" else comment.strip(); is_safe = bool(is_safe) and status == "pending"
    timestamp = now_iso(); actor_id, actor_name = _actor(actor)
    with connect() as db:
        db.execute("BEGIN IMMEDIATE")
        current = db.execute("SELECT * FROM quotes WHERE id=?", (quote_id,)).fetchone()
        if not current: raise KeyError("Quotation not found")
        if current["is_historical"]: raise ValueError("Historical quotations are read-only")
        current_invoices=po_invoices.active(db,"quote_invoices",quote_id)
        invoice_changed=po_invoices.signature(current_invoices)!=po_invoices.signature(normalized_invoices)
        automatic=""
        if status=="po" and (current["status"]!="po" or invoice_changed):
            automatic=po_invoices.automatic_comment(normalized_invoices,"USD",str((actor or {}).get("language") or "en"),current["status"]=="po")
        
        # Blindaje del comentario
        saved_comment=automatic or comment
        changes = {"status":status != current["status"],"loss":loss_reason != current["loss_reason"],
                   "safe":int(is_safe) != int(current["is_safe"]),"invoices":invoice_changed,
                   "po_total":po_total_usd != current["po_total_usd"],"po_date":po_date != current["po_date"],"method":follow_up_type != current["follow_up_type"],
                   "comment":bool(comment)}
        db.execute(
            """UPDATE quotes SET status=?,loss_reason=?,follow_up_type=?,is_safe=?,po_total_usd=?,po_date=?,
            po_detected=CASE WHEN ?='po' THEN 1 ELSE po_detected END,comment=CASE WHEN ?<>'' THEN ? ELSE comment END,
            client_response=?,updated_at=?,last_reviewed_at=?,last_reviewed_by=?,status_changed_at=CASE WHEN status<>? THEN ? ELSE status_changed_at END
            WHERE id=?""",
            (status,loss_reason,follow_up_type,int(is_safe),po_total_usd,po_date,status,saved_comment,saved_comment,client_response,timestamp,timestamp,
             actor_id,status,timestamp,quote_id),
        )
        if status=="po":
            po_invoices.replace(db,"quote_invoices",quote_id,normalized_invoices,actor_id,actor_name,timestamp)
        else:
            po_invoices.clear(db,"quote_invoices",quote_id,timestamp)
        db.execute(
            """INSERT INTO quote_events(quote_id,event_type,from_status,to_status,follow_up_type,note,user_id,user_name,created_at)
            VALUES(?,'review_saved',?,?,?,?,?,?,?)""",
            (quote_id,current["status"],status,follow_up_type,"Review saved",actor_id,actor_name,timestamp),
        )
        if changes["status"]:
            db.execute("""INSERT INTO quote_events(quote_id,event_type,from_status,to_status,follow_up_type,note,user_id,user_name,created_at)
                       VALUES(?,'status_changed',?,?,?,?,?,?,?)""",
                       (quote_id,current["status"],status,follow_up_type,loss_reason if status=="lost" else "",actor_id,actor_name,timestamp))
        if saved_comment:
            db.execute("INSERT INTO quote_comments(quote_id,body,user_id,user_name,created_at) VALUES(?,?,?,?,?)", (quote_id,saved_comment,actor_id,actor_name,timestamp))
            event_type="po_invoices_updated" if automatic and current["status"]=="po" else "po_invoices_registered" if automatic else "comment_added"
            db.execute("INSERT INTO quote_events(quote_id,event_type,follow_up_type,note,user_id,user_name,created_at) VALUES(?,?,?,?,?,?,?)", (quote_id,event_type,follow_up_type,saved_comment,actor_id,actor_name,timestamp))
        if current_invoices and status!="po":
            db.execute("INSERT INTO quote_events(quote_id,event_type,follow_up_type,note,user_id,user_name,created_at) VALUES(?,'po_invoices_cleared',?,'Active PO invoices removed after status change',?,?,?)",(quote_id,follow_up_type,actor_id,actor_name,timestamp))
        for changed,event_type,note in (
            (changes["safe"],"safe_changed","Marked as Safe" if is_safe else "Removed from Safe"),
        ):
            if changed:
                db.execute("INSERT INTO quote_events(quote_id,event_type,follow_up_type,note,user_id,user_name,created_at) VALUES(?,?,?,?,?,?,?)", (quote_id,event_type,follow_up_type,note,actor_id,actor_name,timestamp))
    return get_quote(quote_id)

def dashboard(agent: str = "", start: str = "", end: str = "") -> dict[str, Any]:
    start, end = validate_period(start, end)
    agent_clause = " AND q.nt_agent=?" if agent else ""
    agent_params: list[Any] = [agent] if agent else []
    quote_period = " AND q.quote_date BETWEEN ? AND ?" if start else ""
    quote_params = [*agent_params, *([start, end] if start else [])]
    po_period = " AND q.po_date BETWEEN ? AND ?" if start else ""
    po_params = [*agent_params, *([start, end] if start else [])]
    today = today_local().isoformat()
    with connect() as db:
        pending = dict(db.execute(
            f"""SELECT COUNT(*) AS all_pending,
            SUM(CASE WHEN q.is_safe=0 THEN 1 ELSE 0 END) AS pending,
            SUM(CASE WHEN q.is_safe=1 THEN 1 ELSE 0 END) AS safe,
            COALESCE(SUM(q.total_usd),0) AS pending_value
            FROM quotes q WHERE q.status='pending' AND q.is_archived=0 AND q.is_historical=0
            {agent_clause}{quote_period}""", quote_params).fetchone())
        po = dict(db.execute(
            f"""SELECT COUNT(*) AS po,COALESCE(SUM(q.total_usd),0) AS po_quoted_value,
            COALESCE(SUM(COALESCE(q.po_total_usd,0)),0) AS po_value
            FROM quotes q WHERE q.status='po' AND q.is_archived=0 AND q.is_historical=0
            {agent_clause}{po_period}""", po_params).fetchone())
        lost = db.execute(
            f"""SELECT COUNT(*) AS value FROM quotes q WHERE q.status='lost' AND q.is_archived=0
            AND q.is_historical=0 {agent_clause}{quote_period}""", quote_params).fetchone()["value"]
        pending = {key: (value or 0) for key, value in pending.items()}
        po = {key: (value or 0) for key, value in po.items()}
        counts = {**pending, **po, "lost": lost or 0, "total": int(pending["all_pending"])+int(po["po"])+int(lost or 0)}
        
        # --- NUEVO: Gráfica de Efectividad de Contacto ---
        responses_raw = rows_to_dicts(db.execute(
            f"""SELECT q.client_response, COUNT(*) AS count
            FROM quotes q WHERE q.status='pending' AND q.is_archived=0 AND q.is_historical=0
            {agent_clause}{quote_period} GROUP BY q.client_response""", quote_params).fetchall())
        responses = {"yes": 0, "no": 0, "pending": 0}
        for r in responses_raw:
            val = r.get("client_response") or "pending"
            if val in responses:
                responses[val] += int(r["count"])
        # -------------------------------------------------
        
        priorities = rows_to_dicts(db.execute(
            f"""SELECT q.priority,COUNT(*) AS count,COALESCE(SUM(q.total_usd),0) AS value
            FROM quotes q WHERE q.status='pending' AND q.is_archived=0 AND q.is_historical=0
            {agent_clause}{quote_period} GROUP BY q.priority""", quote_params).fetchall())
        added_today = db.execute(
            f"SELECT COUNT(*) AS value FROM quotes q WHERE date(q.discovered_at)=? AND q.is_historical=0 {agent_clause}",
            (today,*agent_params),
        ).fetchone()["value"]
        po_today = db.execute(
            f"""SELECT COUNT(DISTINCT e.quote_id) AS value FROM quote_events e
            JOIN quotes q ON q.id=e.quote_id WHERE substr(e.created_at,1,10)=? AND e.event_type='status_changed'
            AND e.to_status='po' AND q.is_historical=0 AND q.is_archived=0 {agent_clause}""",
            (today,*agent_params),
        ).fetchone()["value"]
        po_missing = db.execute(
            f"SELECT COUNT(*) AS value FROM quotes q WHERE q.po_detected=1 AND q.po_date IS NULL AND q.status='pending' AND q.is_historical=0 {agent_clause}",
            agent_params,
        ).fetchone()["value"]
        last_import = db.execute("SELECT * FROM imports WHERE workspace='standard' AND status='confirmed' ORDER BY id DESC LIMIT 1").fetchone()
        pending_rows = rows_to_dicts(db.execute(
            f"""SELECT q.* FROM quotes q WHERE q.status='pending' AND q.is_safe=0 AND q.is_historical=0
            AND q.is_archived=0 {agent_clause}""", agent_params).fetchall())
        overdue_all = [decorate_quote(row) for row in pending_rows]
        overdue_all = sorted((row for row in overdue_all if row["overdue"]), key=lambda row: row["days_since_activity"], reverse=True)
    return {"currency":"USD","counts":counts,"responses":responses,"added_today":added_today,"po_today":po_today,"po_missing":po_missing,
            "priorities":priorities,"last_import":dict(last_import) if last_import else None,
            "overdue_count":len(overdue_all),"overdue_quotes":overdue_all[:8],"start_date":start,"end_date":end}

def agents() -> list[str]:
    with connect() as db:
        rows = db.execute("SELECT DISTINCT nt_agent FROM quotes WHERE trim(nt_agent)<>'' ORDER BY nt_agent").fetchall()
    return [row["nt_agent"] for row in rows]


def executive_pipeline(rows: list[dict[str, Any]], workspace: str) -> dict[str, Any]:
    """Summarize the current pending pipeline for the one-page report."""
    amount_key = "unit_price" if workspace == "special" else "total_usd"
    decorated = [decorate_quote(dict(row)) for row in rows]
    priorities = []
    for priority in "SABC":
        matching = [row for row in decorated if row.get("priority") == priority]
        priorities.append({
            "priority": priority,
            "count": len(matching),
            "value": sum(float(row.get(amount_key) or 0) for row in matching),
        })

    followup_soon = [row for row in decorated if not row.get("is_safe") and 9 <= row["days_since_activity"] <= 14]
    followup_due = [row for row in decorated if not row.get("is_safe") and row["days_since_activity"] >= 15]
    stale = [row for row in decorated if row["quote_age_days"] > 90]
    po_missing = [
        row for row in decorated
        if workspace == "standard" and row.get("po_detected") and not row.get("po_date")
    ]

    def alert(rows_for_alert: list[dict[str, Any]]) -> dict[str, Any]:
        return {
            "count": len(rows_for_alert),
            "value": sum(float(row.get(amount_key) or 0) for row in rows_for_alert),
        }

    candidates = []
    for row in decorated:
        categories = []
        if row.get("priority") in {"S", "A"} and not row.get("is_safe") and row["days_since_activity"] >= 15:
            categories.append(0)
        if row["quote_age_days"] > 90:
            categories.append(1)
        if workspace == "standard" and row.get("po_detected") and not row.get("po_date"):
            categories.append(2)
        if not categories:
            continue
        candidates.append((min(categories), -float(row.get(amount_key) or 0), -row["days_since_activity"], row))
    candidates.sort(key=lambda item: item[:3])
    action_rows = []
    for _, _, _, row in candidates[:10]:
        action_rows.append({
            "priority": row.get("priority") or "C",
            "folio": row.get("folio") or row.get("source_quote_number") or f"SPQ-{row.get('id',0):05d}",
            "distributor_code": row.get("distributor_code") or "—",
            "end_user": row.get("end_user") or row.get("customer_name") or row.get("receptor") or "—",
            "amount": float(row.get(amount_key) or 0),
            "days_since_activity": int(row.get("days_since_activity") or 0),
            "nt_agent": row.get("nt_agent") or "—",
        })
    return {
        "pending_count": len(decorated),
        "priorities": priorities,
        "alerts": {
            "followup_9_14": alert(followup_soon),
            "followup_15_plus": alert(followup_due),
            "stale_90_plus": alert(stale),
            "po_date_missing": alert(po_missing),
        },
        "action_rows": action_rows,
        "action_remaining": max(0, len(candidates) - len(action_rows)),
    }


def report_activity(start_date: str, end_date: str, actor_user_id: int | None = None, agent: str = "") -> dict[str, Any]:
    start = date.fromisoformat(start_date); end = date.fromisoformat(end_date)
    if end < start: raise ValueError("End date cannot be before start date")
    event_actor = " AND e.user_id=?" if actor_user_id else ""; event_agent = " AND q.nt_agent=?" if agent else ""
    event_params: list[Any] = [start_date,end_date] + ([actor_user_id] if actor_user_id else []) + ([agent] if agent else [])
    quote_agent = " AND q.nt_agent=?" if agent else ""; quote_params: list[Any] = [start_date,end_date] + ([agent] if agent else [])
    created_actor = " AND q.created_by_user_id=?" if actor_user_id else ""
    created_params: list[Any] = [start_date,end_date] + ([actor_user_id] if actor_user_id else []) + ([agent] if agent else [])
    with connect() as db:
        events = rows_to_dicts(db.execute(
            f"""SELECT e.*,q.folio,q.quote_date,q.distributor_code,q.end_user,q.nt_agent,q.priority,q.status,q.total_usd,q.net_total_usd,
            q.po_total_usd,q.po_date,q.loss_reason,
            (SELECT COUNT(*) FROM quote_invoices i WHERE i.quote_id=q.id AND i.active=1) AS invoice_count
            FROM quote_events e JOIN quotes q ON q.id=e.quote_id
            WHERE date(e.created_at) BETWEEN ? AND ? {event_actor} {event_agent} AND q.is_historical=0 AND q.is_archived=0
            ORDER BY e.created_at,e.id""", event_params).fetchall())
        new_quotes = rows_to_dicts(db.execute(
            f"SELECT q.* FROM quotes q WHERE date(q.discovered_at) BETWEEN ? AND ? {created_actor} {quote_agent} AND q.is_historical=0 AND q.is_archived=0 ORDER BY q.discovered_at", created_params).fetchall())
        pending_rows = rows_to_dicts(db.execute(
            f"SELECT q.* FROM quotes q WHERE q.status='pending' AND q.quote_date BETWEEN ? AND ? {quote_agent} AND q.is_historical=0 AND q.is_archived=0",
            quote_params,
        ).fetchall())
    reviews = [e for e in events if e["event_type"]=="review_saved"]
    status_events = [e for e in events if e["event_type"]=="status_changed"]
    lost_by_quote = {e["quote_id"]:e for e in status_events if e["to_status"]=="lost"}
    po_by_quote = {e["quote_id"]:e for e in status_events if e["to_status"]=="po"}
    loss_breakdown: dict[str,int] = {}; loss_breakdown_detail: dict[str,dict[str,Any]] = {}
    for e in lost_by_quote.values():
        reason = e.get("note") or e.get("loss_reason") or "Not specified"; loss_breakdown[reason] = loss_breakdown.get(reason,0)+1
        detail=loss_breakdown_detail.setdefault(reason,{"count":0,"value":0.0}); detail["count"]+=1; detail["value"]+=float(e.get("total_usd") or 0)
    reviewed_ids = {e["quote_id"] for e in reviews}; reviewed = []
    for quote_id in reviewed_ids:
        matching = [e for e in events if e["quote_id"]==quote_id and e["event_type"] in {"review_saved","comment_added","status_changed","safe_changed","po_invoices_registered","po_invoices_updated","po_invoices_cleared"}]
        base = matching[0]
        reviewed.append({"quote_id":quote_id,"folio":base["folio"],"receptor":base["end_user"],"nt_agent":base["nt_agent"],
                         "priority":base["priority"],"status":base["status"],"total":base["total_usd"],
                         "net_total":base["net_total_usd"],"po_total":base["po_total_usd"],"events":matching})
    po_rows = [{"folio":e["folio"],"receptor":e["end_user"],"quote_date":e["quote_date"],"po_date":e["po_date"],"quoted_total":e["total_usd"],
                "po_total":e["po_total_usd"],"invoice_count":e.get("invoice_count",0),
                "variance":float(e["po_total_usd"] or 0)-float(e["total_usd"] or 0)} for e in po_by_quote.values()]
    lost_rows = [{"folio":e["folio"],"receptor":e["end_user"],"quoted_total":e["total_usd"],"reason":e.get("note") or e.get("loss_reason") or "Not specified","date":e["created_at"][:10]} for e in lost_by_quote.values()]
    pipeline=executive_pipeline(pending_rows,"standard")
    cycle_days=[]
    for row in po_rows:
        try:
            elapsed=(date.fromisoformat(str(row["po_date"])[:10])-date.fromisoformat(str(row["quote_date"])[:10])).days
            if elapsed>=0: cycle_days.append(elapsed)
        except (TypeError,ValueError):
            pass
    resolved=len(po_rows)+len(lost_rows)
    return {"workspace":"standard","currency":"USD","start_date":start_date,"end_date":end_date,
            "new_quotes":len(new_quotes),"quotes_reviewed":len(reviewed_ids),"status_changes":len(status_events),
            "po_changes":len(po_rows),"lost_changes":len(lost_rows),"pending_value":sum(float(row.get("total_usd") or 0) for row in pending_rows),
            "pending_count":pipeline["pending_count"],"new_value":sum(float(row.get("total_usd") or 0) for row in new_quotes),
            "lost_value":sum(float(row.get("quoted_total") or 0) for row in lost_rows),
            "conversion_rate":(len(po_rows)/resolved*100) if resolved else None,
            "average_po_days":(sum(cycle_days)/len(cycle_days)) if cycle_days else None,
            "po_quoted_value":sum(float(r["quoted_total"] or 0) for r in po_rows),"po_value":sum(float(r["po_total"] or 0) for r in po_rows),
            "loss_breakdown":loss_breakdown,"loss_breakdown_detail":loss_breakdown_detail,"po_rows":po_rows,"lost_rows":lost_rows,
            "reviewed":reviewed,"new_rows":new_quotes,**{key:value for key,value in pipeline.items() if key!="pending_count"}}


def monthly_report(months: int = 12, agent: str = "") -> list[dict[str, Any]]:
    clause = " AND nt_agent=?" if agent else ""; params: tuple[Any,...] = (agent,) if agent else ()
    with connect() as db:
        rows = db.execute(
            f"""SELECT substr(quote_date,1,7) AS month,COUNT(*) AS quotes,SUM(status='pending') AS pending,
            SUM(status='po') AS po,SUM(status='lost') AS lost,COALESCE(SUM(total_usd),0) AS quoted_value,
            COALESCE(SUM(CASE WHEN status='po' THEN po_total_usd ELSE 0 END),0) AS po_value
            FROM quotes WHERE is_historical=0 AND is_archived=0 {clause} GROUP BY month ORDER BY month DESC LIMIT ?""",
            (*params,months)).fetchall()
    return list(reversed(rows_to_dicts(rows)))


def daily_activity(report_date: str, agent: str = "") -> dict[str, Any]:
    data = report_activity(report_date,report_date,None,agent)
    data.update({"date":report_date,"agent":agent,"report_currency":"USD","report_title":"Daily follow-up report",
                 "activity_count":data["quotes_reviewed"],"quotes":data["reviewed"],"notes":[],"comment_changes":sum(
                     len([e for e in q["events"] if e["event_type"]=="comment_added"]) for q in data["reviewed"]),
                 "reviews_without_changes":0,"follow_up_methods":{m:len({q["quote_id"] for q in data["reviewed"] if any(e["follow_up_type"]==m for e in q["events"])}) for m in ("email","call","visit")}})
    return data


def quote_path(quote_id: int) -> Path:
    with connect() as db:
        row = db.execute("SELECT source_path FROM quotes WHERE id=?", (quote_id,)).fetchone()
    if not row: raise KeyError("Quotation not found")
    return Path(row["source_path"])
