from __future__ import annotations

import hashlib
import math
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

import database
import po_invoices


def initialize() -> None:
    with database.connect() as db:
        db.executescript(
            """
            CREATE TABLE IF NOT EXISTS special_quotes (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                source_key TEXT NOT NULL UNIQUE,
                quote_date TEXT NOT NULL,
                customer_number TEXT NOT NULL,
                customer_name TEXT NOT NULL DEFAULT '',
                code TEXT NOT NULL,
                description TEXT NOT NULL,
                quantity REAL NOT NULL,
                unit_price REAL NOT NULL,
                source_total REAL NOT NULL,
                source_currency TEXT NOT NULL,
                units_per_usd REAL NOT NULL,
                total_usd REAL NOT NULL,
                priority TEXT NOT NULL CHECK (priority IN ('S', 'A', 'B', 'C')),
                nt_agent TEXT NOT NULL DEFAULT '',
                status TEXT NOT NULL CHECK (status IN ('pending', 'po', 'lost')) DEFAULT 'pending',
                follow_up_type TEXT NOT NULL DEFAULT '',
                is_safe INTEGER NOT NULL DEFAULT 0,
                po_total_usd REAL,
                comment TEXT NOT NULL DEFAULT '',
                loss_reason TEXT NOT NULL DEFAULT '',
                source_file TEXT NOT NULL,
                source_sheet TEXT NOT NULL,
                source_row INTEGER NOT NULL,
                discovered_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                status_changed_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS special_quote_events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                quote_id INTEGER NOT NULL REFERENCES special_quotes(id) ON DELETE CASCADE,
                event_type TEXT NOT NULL, from_status TEXT, to_status TEXT,
                follow_up_type TEXT NOT NULL DEFAULT '', note TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS special_quote_comments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                quote_id INTEGER NOT NULL REFERENCES special_quotes(id) ON DELETE CASCADE,
                body TEXT NOT NULL,user_id INTEGER,user_name TEXT NOT NULL DEFAULT 'Historical data',
                created_at TEXT NOT NULL,is_legacy INTEGER NOT NULL DEFAULT 0
            );
            CREATE TABLE IF NOT EXISTS special_quote_invoices (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                quote_id INTEGER NOT NULL REFERENCES special_quotes(id) ON DELETE CASCADE,
                invoice_date TEXT NOT NULL,invoice_series TEXT NOT NULL,invoice_number TEXT NOT NULL,
                amount REAL NOT NULL CHECK(amount>0),currency TEXT NOT NULL DEFAULT 'JPY',
                created_by_user_id INTEGER,created_by_name TEXT NOT NULL DEFAULT '',created_at TEXT NOT NULL,
                active INTEGER NOT NULL DEFAULT 1,is_legacy INTEGER NOT NULL DEFAULT 0,superseded_at TEXT
            );
            CREATE TABLE IF NOT EXISTS special_imports (
                id INTEGER PRIMARY KEY AUTOINCREMENT,filename TEXT NOT NULL,source_currency TEXT NOT NULL,
                units_per_usd REAL NOT NULL,rows_seen INTEGER NOT NULL,inserted INTEGER NOT NULL,
                updated INTEGER NOT NULL,errors INTEGER NOT NULL,error_detail TEXT NOT NULL DEFAULT '',imported_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_special_status ON special_quotes(status);
            CREATE INDEX IF NOT EXISTS idx_special_agent ON special_quotes(nt_agent);
            CREATE INDEX IF NOT EXISTS idx_special_date ON special_quotes(quote_date);
            CREATE INDEX IF NOT EXISTS idx_special_events_created ON special_quote_events(created_at);
            CREATE INDEX IF NOT EXISTS idx_special_invoices_active ON special_quote_invoices(quote_id,active);
            """
        )
        db.execute("""
            CREATE TABLE IF NOT EXISTS po_invoices (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                quote_id INTEGER NOT NULL,
                invoice_date TEXT,
                invoice_series TEXT,
                invoice_number TEXT,
                invoice_amount REAL,
                FOREIGN KEY(quote_id) REFERENCES quotes(id)
            )
        """)
        database._ensure_columns(db, "special_quotes", {
            "source_quote_number":"TEXT NOT NULL DEFAULT ''","rank":"TEXT NOT NULL DEFAULT ''","po_date":"TEXT",
            "is_archived":"INTEGER NOT NULL DEFAULT 0","archive_reason":"TEXT NOT NULL DEFAULT ''","import_id":"INTEGER",
            "created_by_user_id":"INTEGER","created_by_name":"TEXT NOT NULL DEFAULT ''","last_reviewed_at":"TEXT",
            "last_reviewed_by":"INTEGER","client_response":"TEXT NOT NULL DEFAULT 'pending'",
        })
        database._ensure_columns(db, "special_quote_events", {
            "user_id":"INTEGER","user_name":"TEXT NOT NULL DEFAULT 'Historical data'",
        })
        db.execute(
            """INSERT INTO special_quote_comments(quote_id,body,user_name,created_at,is_legacy)
            SELECT q.id,q.comment,'Historical data',q.updated_at,1 FROM special_quotes q
            WHERE trim(q.comment)<>'' AND NOT EXISTS(
                SELECT 1 FROM special_quote_comments c WHERE c.quote_id=q.id AND c.is_legacy=1)"""
        )
        database.backfill_review_metadata(db,"special_quotes","special_quote_events","special_quote_comments")
        _reconcile_duplicates(db)
        po_invoices.migrate_legacy(db,"special_quotes","special_quote_invoices","special_quote_events","JPY")


def _text(value: Any) -> str:
    if value is None: return ""
    if isinstance(value, float) and value.is_integer(): return str(int(value))
    return str(value).strip()


def _date(value: Any) -> date:
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    text = _text(value)
    for pattern in ("%Y-%m-%d","%Y/%m/%d","%d/%m/%Y","%m/%d/%Y"):
        try: return datetime.strptime(text, pattern).date()
        except ValueError: continue
    raise ValueError(f"Unrecognized date: {text or 'blank'}")


def _number(value: Any, label: str) -> float:
    try: result = float(str(value).replace(",", "").replace("¥", "").strip())
    except (TypeError, ValueError) as exc: raise ValueError(f"{label} must be numeric") from exc
    if result < 0: raise ValueError(f"{label} cannot be negative")
    return result


def _source_key(record: dict[str, Any]) -> str:
    quote_number = record["source_quote_number"].strip().casefold()
    if quote_number:
        raw = f"quote|{quote_number}"
    else:
        raw = "|".join([
            "legacy",record["quote_date"],record["code"].strip().casefold(),
            record["description"].strip().casefold(),
        ])
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _row_source_key(row: Any) -> str:
    return _source_key({
        "source_quote_number":_text(row["source_quote_number"]),"quote_date":_text(row["quote_date"]),
        "code":_text(row["code"]),"description":_text(row["description"]),
    })


def _merge_duplicate(db: Any, canonical: Any, duplicate: Any) -> None:
    """Keep one active quotation while preserving every follow-up record."""
    canonical_id, duplicate_id = int(canonical["id"]), int(duplicate["id"])
    canonical_review = canonical["last_reviewed_at"] or ""
    duplicate_review = duplicate["last_reviewed_at"] or ""
    if duplicate_review and duplicate_review > canonical_review:
        db.execute(
            """UPDATE special_quotes SET status=?,follow_up_type=?,is_safe=?,po_total_usd=?,po_date=?,comment=?,
            loss_reason=?,last_reviewed_at=?,last_reviewed_by=?,status_changed_at=? WHERE id=?""",
            (duplicate["status"],duplicate["follow_up_type"],duplicate["is_safe"],duplicate["po_total_usd"],
             duplicate["po_date"],duplicate["comment"],duplicate["loss_reason"],duplicate["last_reviewed_at"],
             duplicate["last_reviewed_by"],duplicate["status_changed_at"],canonical_id),
        )
    db.execute("UPDATE special_quote_comments SET quote_id=? WHERE quote_id=?",(canonical_id,duplicate_id))
    db.execute("UPDATE special_quote_events SET quote_id=? WHERE quote_id=?",(canonical_id,duplicate_id))
    db.execute("UPDATE special_quote_invoices SET quote_id=? WHERE quote_id=?",(canonical_id,duplicate_id))
    earliest=min(str(canonical["discovered_at"]),str(duplicate["discovered_at"]))
    db.execute("UPDATE special_quotes SET discovered_at=? WHERE id=?",(earliest,canonical_id))
    db.execute(
        "UPDATE special_quotes SET is_archived=1,archive_reason=? WHERE id=?",
        (f"Merged duplicate into quotation record {canonical_id}",duplicate_id),
    )
    db.execute(
        """INSERT INTO special_quote_events(quote_id,event_type,note,user_name,created_at)
        VALUES(?,'duplicate_merged',?,'System',?)""",
        (canonical_id,f"Merged duplicate record {duplicate_id}",database.now_iso()),
    )


def _reconcile_duplicates(db: Any) -> int:
    """Merge duplicates created by older import formats and normalize active keys."""
    merged = 0
    duplicate_numbers = db.execute(
        """SELECT source_quote_number FROM special_quotes
        WHERE is_archived=0 AND trim(source_quote_number)<>''
        GROUP BY source_quote_number HAVING COUNT(*)>1"""
    ).fetchall()
    for group in duplicate_numbers:
        rows=db.execute(
            """SELECT * FROM special_quotes WHERE is_archived=0 AND source_quote_number=?
            ORDER BY CASE WHEN import_id IS NOT NULL THEN 0 ELSE 1 END,updated_at DESC,id DESC""",
            (group["source_quote_number"],),
        ).fetchall()
        canonical=rows[0]
        for duplicate in rows[1:]:
            _merge_duplicate(db,canonical,duplicate); merged+=1
            canonical=db.execute("SELECT * FROM special_quotes WHERE id=?",(canonical["id"],)).fetchone()

    legacy_rows=db.execute(
        "SELECT * FROM special_quotes WHERE is_archived=0 AND trim(source_quote_number)='' ORDER BY id"
    ).fetchall()
    for legacy in legacy_rows:
        candidates=db.execute(
            """SELECT * FROM special_quotes WHERE is_archived=0 AND trim(source_quote_number)<>''
            AND quote_date=? AND lower(trim(code))=lower(trim(?)) AND lower(trim(description))=lower(trim(?))""",
            (legacy["quote_date"],legacy["code"],legacy["description"]),
        ).fetchall()
        if len(candidates)>1:
            candidates=[row for row in candidates if float(row["quantity"])==float(legacy["quantity"])
                        and float(row["unit_price"])==float(legacy["unit_price"])]
        if len(candidates)==1:
            _merge_duplicate(db,candidates[0],legacy); merged+=1

    active=db.execute("SELECT * FROM special_quotes WHERE is_archived=0").fetchall()
    for row in active:
        stable_key=_row_source_key(row)
        conflict=db.execute("SELECT id FROM special_quotes WHERE source_key=? AND id<>?",(stable_key,row["id"])).fetchone()
        if not conflict:
            db.execute("UPDATE special_quotes SET source_key=? WHERE id=?",(stable_key,row["id"]))
    return merged


def _find_existing(db: Any, record: dict[str, Any]) -> Any:
    quote_number=record["source_quote_number"].strip()
    if quote_number:
        current=db.execute(
            """SELECT * FROM special_quotes WHERE source_quote_number=?
            ORDER BY is_archived,updated_at DESC,id DESC LIMIT 1""",(quote_number,)
        ).fetchone()
        if current: return current
    current=db.execute("SELECT * FROM special_quotes WHERE source_key=?",(record["source_key"],)).fetchone()
    if current: return current
    candidates=db.execute(
        """SELECT * FROM special_quotes WHERE trim(source_quote_number)='' AND
        quote_date=? AND lower(trim(code))=lower(trim(?)) AND lower(trim(description))=lower(trim(?))
        AND archive_reason NOT LIKE 'Merged duplicate%' ORDER BY is_archived,id""",
        (record["quote_date"],record["code"],record["description"]),
    ).fetchall()
    if len(candidates)==1: return candidates[0]
    exact=[row for row in candidates if float(row["quantity"])==float(record["quantity"])
           and float(row["unit_price"])==float(record["unit_price"])]
    return exact[0] if len(exact)==1 else None


def _business_changed(current: Any, record: dict[str, Any]) -> bool:
    text_fields=("source_quote_number","quote_date","customer_name","rank","code","description","priority","nt_agent","archive_reason")
    if any(_text(current[field])!=_text(record[field]) for field in text_fields): return True
    number_fields=("quantity","unit_price")
    if any(float(current[field])!=float(record[field]) for field in number_fields): return True
    return int(current["is_archived"])!=int(record["is_archived"])


def _apply_visibility(records: list[dict[str, Any]]) -> None:
    groups: dict[str,list[dict[str,Any]]] = {}
    for record in records:
        groups.setdefault(record["description"].strip().casefold(), []).append(record)
    for group in groups.values():
        if len(group) == 1:
            group[0]["is_archived"] = 0; group[0]["archive_reason"] = ""
            continue
        by_date: dict[str,list[dict[str,Any]]] = {}
        for record in group: by_date.setdefault(record["quote_date"], []).append(record)
        for same_date in by_date.values():
            candidates = [record for record in same_date if float(record["quantity"]) == 1]
            winner = sorted(candidates, key=lambda item:(-item["unit_price"],item["source_row"]))[0] if candidates else None
            for record in same_date:
                record["is_archived"] = int(record is not winner)
                record["archive_reason"] = "Duplicate model/quantity rule" if record is not winner else ""


def _assign_priorities(records: list[dict[str, Any]]) -> None:
    visible = [record for record in records if not record["is_archived"]]
    total = len(visible)
    if not total: return
    s_cutoff = math.ceil(total*.10); a_cutoff = math.ceil(total*.30); b_cutoff = math.ceil(total*.60)
    ranked = sorted(visible, key=lambda item:(-item["unit_price"],item["source_row"]))
    for index,record in enumerate(ranked,1):
        record["priority"] = "S" if index<=s_cutoff else "A" if index<=a_cutoff else "B" if index<=b_cutoff else "C"
    for record in records:
        if record["is_archived"]: record["priority"] = "C"


def parse_workbook(content: bytes, filename: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if not filename.lower().endswith(".xlsx"): raise ValueError("Upload an .xlsx Excel file")
    if len(content)>25*1024*1024: raise ValueError("The Excel file cannot exceed 25 MB")
    try: workbook = load_workbook(BytesIO(content),data_only=True,read_only=True)
    except Exception as exc: raise ValueError("The Excel workbook could not be read") from exc
    worksheet = workbook.active
    worksheet_rows = list(worksheet.iter_rows(min_col=1, max_col=27, values_only=True))
    header_row = 1
    for row,row_values in enumerate(worksheet_rows[:10],start=1):
        if _text(row_values[1]).casefold() in {"quote date","date","fecha"}:
            header_row = row; break
    records: list[dict[str,Any]] = []; errors: list[dict[str,Any]] = []
    for row,row_values in enumerate(worksheet_rows[header_row:],start=header_row+1):
        values = {col:row_values[col-1] for col in (1,2,5,6,7,8,9,10,27)}
        if not any(value not in (None,"") for value in values.values()): continue
        try:
            model = _text(values[8]); company = _text(values[5]); code = _text(values[7])
            if not model or not company or not code: raise ValueError("Company, item code, and model are required")
            record = {
                "source_quote_number":_text(values[1]),"quote_date":_date(values[2]).isoformat(),
                "customer_number":"","customer_name":company,"rank":_text(values[6]),"code":code,
                "description":model,"quantity":_number(values[9],"Quantity"),"unit_price":_number(values[10],"Unit price"),
                "source_total":_number(values[10],"Unit price"),"source_currency":"JPY","units_per_usd":1,
                "total_usd":_number(values[10],"Unit price"),"nt_agent":_text(values[27]) or "Unassigned",
                "source_file":filename,"source_sheet":worksheet.title,"source_row":row,"priority":"C",
                "is_archived":0,"archive_reason":"",
            }
            record["source_key"] = _source_key(record); records.append(record)
        except Exception as exc: errors.append({"row":row,"error":str(exc)})
    if not records and not errors: raise ValueError("The workbook does not contain special quotation rows")
    _apply_visibility(records); _assign_priorities(records)
    return records,errors


def preview_workbook(content: bytes, filename: str) -> dict[str,Any]:
    records,errors = parse_workbook(content,filename)
    unique = {record["source_key"]:record for record in records}
    duplicates = len(records)-len(unique)
    with database.connect() as db:
        new=updated=unchanged=0
        for record in unique.values():
            current=_find_existing(db,record)
            if not current: new+=1
            elif _business_changed(current,record): updated+=1
            else: unchanged+=1
    return {"rows_seen":len(records),"valid":len(unique),"new":new,"updated":updated,"unchanged":unchanged,
            "duplicates":duplicates,"archived":sum(record["is_archived"] for record in unique.values()),"errors":len(errors),
            "error_detail":errors[:20],"sample":[record for record in unique.values() if not record["is_archived"]][:8]}


def import_workbook(content: bytes, filename: str, actor: dict[str,Any]|None=None, import_id: int|None=None) -> dict[str,Any]:
    records,errors = parse_workbook(content,filename); unique={record["source_key"]:record for record in records}
    actor_id,actor_name=database._actor(actor); timestamp=database.now_iso(); inserted=updated=unchanged=0
    with database.connect() as db:
        db.execute("BEGIN IMMEDIATE")
        merged=_reconcile_duplicates(db)
        for record in unique.values():
            current=_find_existing(db,record)
            values=(record["source_quote_number"],record["quote_date"],record["customer_name"],record["rank"],record["code"],
                    record["description"],record["quantity"],record["unit_price"],record["unit_price"],record["unit_price"],
                    record["priority"],record["nt_agent"],record["source_file"],record["source_sheet"],record["source_row"],
                    record["is_archived"],record["archive_reason"],import_id)
            if current:
                if not _business_changed(current,record):
                    unchanged+=1; continue
                db.execute("""UPDATE special_quotes SET source_quote_number=?,quote_date=?,customer_name=?,rank=?,code=?,description=?,
                    quantity=?,unit_price=?,source_total=?,total_usd=?,priority=?,nt_agent=?,source_file=?,source_sheet=?,source_row=?,
                    is_archived=?,archive_reason=?,import_id=?,source_key=?,source_currency='JPY',units_per_usd=1,updated_at=? WHERE id=?""",
                    (*values,record["source_key"],timestamp,current["id"]))
                db.execute(
                    """INSERT INTO special_quote_events(quote_id,event_type,note,user_id,user_name,created_at)
                    VALUES(?,'source_updated','Quotation data updated from Excel',?,?,?)""",
                    (current["id"],actor_id,actor_name,timestamp),
                )
                updated+=1
            else:
                cursor=db.execute("""INSERT INTO special_quotes(source_key,source_quote_number,quote_date,customer_number,customer_name,
                    rank,code,description,quantity,unit_price,source_total,source_currency,units_per_usd,total_usd,priority,nt_agent,
                    status,source_file,source_sheet,source_row,is_archived,archive_reason,import_id,created_by_user_id,created_by_name,
                    discovered_at,updated_at,status_changed_at)
                    VALUES(?,?,?,'',?,?,?,?,?,?,?,'JPY',1,?,?,?,'pending',?,?,?,?,?,?,?,?,?,?,?)""",
                    (record["source_key"],record["source_quote_number"],record["quote_date"],record["customer_name"],record["rank"],
                     record["code"],record["description"],record["quantity"],record["unit_price"],record["unit_price"],record["unit_price"],
                     record["priority"],record["nt_agent"],record["source_file"],record["source_sheet"],record["source_row"],
                     record["is_archived"],record["archive_reason"],import_id,actor_id,actor_name,timestamp,timestamp,timestamp))
                db.execute("INSERT INTO special_quote_events(quote_id,event_type,to_status,note,user_id,user_name,created_at) VALUES(?,'created','pending','Imported from Excel',?,?,?)",(cursor.lastrowid,actor_id,actor_name,timestamp))
                inserted+=1
        merged+=_reconcile_duplicates(db)
    return {"rows_seen":len(records),"inserted":inserted,"updated":updated,"unchanged":unchanged,
            "duplicates":len(records)-len(unique),"merged":merged,
            "archived":sum(r["is_archived"] for r in unique.values()),"errors":len(errors),"error_detail":errors[:20]}


def _decorate(row: dict[str,Any]) -> dict[str,Any]:
    row["folio"] = row.get("source_quote_number") or f"SPQ-{int(row['id']):05d}"
    row["receptor"] = row.get("customer_name") or "Unassigned customer"
    row["currency"] = "JPY"; row["quoted_amount"] = row.get("unit_price",0); row["po_total"] = row.get("po_total_usd")
    return database.decorate_quote(row)


def list_quotes(filters: dict[str,str]) -> list[dict[str,Any]]:
    clauses=["q.is_archived=?"]; values:[Any]=[1 if filters.get("archive")=="1" else 0]
    if filters.get("status"): clauses.append("q.status=?"); values.append(filters["status"])
    if filters.get("agent"): clauses.append("q.nt_agent=?"); values.append(filters["agent"])
    if filters.get("priority"): clauses.append("q.priority=?"); values.append(filters["priority"])
    if filters.get("rank"): clauses.append("q.rank=?"); values.append(filters["rank"])
    if filters.get("end_user"): clauses.append("q.customer_name LIKE ?"); values.append(f"%{filters['end_user']}%")
    if filters.get("safe") in {"0","1"}: clauses.append("q.is_safe=?"); values.append(int(filters["safe"]))
    start,end=database.validate_period(filters.get("start",""),filters.get("end",""))
    if start: clauses.append("q.quote_date BETWEEN ? AND ?"); values.extend([start,end])
    if filters.get("search"):
        needle=f"%{filters['search']}%"; clauses.append("(q.source_quote_number LIKE ? OR q.customer_name LIKE ? OR q.code LIKE ? OR q.description LIKE ? OR q.rank LIKE ?)"); values.extend([needle]*5)
    orders={"priority":"CASE q.priority WHEN 'S' THEN 1 WHEN 'A' THEN 2 WHEN 'B' THEN 3 ELSE 4 END,q.unit_price DESC",
            "date_desc":"q.quote_date DESC,q.id DESC","date_asc":"q.quote_date ASC,q.id ASC","folio_desc":"q.source_quote_number DESC",
            "folio_asc":"q.source_quote_number ASC","amount_desc":"q.unit_price DESC","amount_asc":"q.unit_price ASC"}
    order=orders.get(filters.get("order","priority"),orders["priority"])
    with database.connect() as db: rows=database.rows_to_dicts(db.execute(f"SELECT q.* FROM special_quotes q WHERE {' AND '.join(clauses)} ORDER BY {order}",values).fetchall())
    return [_decorate(row) for row in rows]


def list_managed(filters: dict[str,str]|None=None) -> list[dict[str,Any]]:
    filters=dict(filters or {}); unmanaged=filters.get("status")=="unmanaged"
    if unmanaged: filters["status"]="pending"
    rows=list_quotes(filters); result=[]
    for row in rows:
        if bool(row.get("last_reviewed_at"))==unmanaged: continue
        if filters.get("overdue")=="1" and not row["overdue"]: continue
        result.append(row)
    if filters.get("order","status_activity")=="status_activity":
        status_order={"pending":0,"lost":1,"po":2}
        result.sort(key=lambda row:(status_order.get(row["status"],9),-(datetime.fromisoformat(row.get("last_reviewed_at") or row["discovered_at"]).timestamp())))
    else:
        result.sort(key=lambda row:row.get("last_reviewed_at") or row["discovered_at"],reverse=filters.get("order")!="oldest_activity")
    return result


def list_management(filters:dict[str,str]|None=None)->list[dict[str,Any]]:
    filters=dict(filters or {}); filters.pop("archive",None); unmanaged=filters.get("status")=="unmanaged"
    if unmanaged: filters["status"]="pending"
    rows=list_quotes(filters)
    if unmanaged: rows=[row for row in rows if not row.get("last_reviewed_at")]
    orders=filters.get("order","newest_activity")
    if orders in {"newest_activity","oldest_activity"}:
        rows.sort(key=lambda row:row.get("last_activity_at") or "",reverse=orders=="newest_activity")
    return rows


def get_quote(quote_id:int)->dict[str,Any]:
    with database.connect() as db:
        row=db.execute("SELECT * FROM special_quotes WHERE id=?",(quote_id,)).fetchone()
        if not row: raise KeyError("Quotation not found")
        comments=database.rows_to_dicts(db.execute("SELECT * FROM special_quote_comments WHERE quote_id=? ORDER BY created_at,id",(quote_id,)).fetchall())
        events=database.rows_to_dicts(db.execute("SELECT * FROM special_quote_events WHERE quote_id=? ORDER BY created_at,id",(quote_id,)).fetchall())
        invoices=po_invoices.active(db,"special_quote_invoices",quote_id)
    result=_decorate(dict(row)); result["comments"]=comments; result["events"]=events; result["invoices"]=invoices; result["read_only"]=False
    return result


def update_quote(quote_id:int,status:str,comment:str,loss_reason:str,follow_up_type:str,is_safe:bool,
                 po_total_usd:float|None,po_date:str|None=None,actor:dict[str,Any]|None=None,
                 invoices:list[dict[str,Any]]|None=None,client_response:str="pending")->dict[str,Any]:
    if status not in {"pending","po","lost"}: raise ValueError("Invalid status")
    if follow_up_type not in {"email","call","visit"}: raise ValueError("Select a follow-up method: E-mail, Call, or Visit")
    loss_reason=loss_reason.strip()
    if status=="lost" and loss_reason not in database.LOSS_REASONS: raise ValueError("Select a valid loss reason")
    if status!="lost": loss_reason=""
    if client_response not in {"pending","yes","no"}: raise ValueError("Invalid client response")
    normalized_invoices=po_invoices.normalize(invoices,"JPY",legacy_total=po_total_usd,legacy_date=po_date,
        legacy_number=f"PO-{quote_id}") if status=="po" else []
    if status=="po": po_total_usd,po_date=po_invoices.totals(normalized_invoices)
    else: po_total_usd=None; po_date=None
    comment="" if status=="po" else comment.strip(); is_safe=bool(is_safe) and status=="pending"; timestamp=database.now_iso(); actor_id,actor_name=database._actor(actor)
    with database.connect() as db:
        db.execute("BEGIN IMMEDIATE"); current=db.execute("SELECT * FROM special_quotes WHERE id=?",(quote_id,)).fetchone()
        if not current: raise KeyError("Quotation not found")
        current_invoices=po_invoices.active(db,"special_quote_invoices",quote_id)
        invoice_changed=po_invoices.signature(current_invoices)!=po_invoices.signature(normalized_invoices)
        automatic=""
        if status=="po" and (current["status"]!="po" or invoice_changed):
            automatic=po_invoices.automatic_comment(normalized_invoices,"JPY",str((actor or {}).get("language") or "en"),current["status"]=="po")
        saved_comment=automatic or comment
        changes={"status":status!=current["status"],"loss":loss_reason!=current["loss_reason"],"safe":int(is_safe)!=int(current["is_safe"]),
                 "invoices":invoice_changed,
                 "po_total":po_total_usd!=current["po_total_usd"],"po_date":po_date!=current["po_date"],
                 "method":follow_up_type!=current["follow_up_type"],"comment":bool(comment)}
        db.execute("""UPDATE special_quotes SET status=?,loss_reason=?,follow_up_type=?,is_safe=?,po_total_usd=?,po_date=?,
            comment=CASE WHEN ?<>'' THEN ? ELSE comment END,client_response=?,updated_at=?,last_reviewed_at=?,last_reviewed_by=?,
            status_changed_at=CASE WHEN status<>? THEN ? ELSE status_changed_at END WHERE id=?""",
            (status,loss_reason,follow_up_type,int(is_safe),po_total_usd,po_date,saved_comment,saved_comment,client_response,timestamp,timestamp,actor_id,status,timestamp,quote_id))
        if status=="po": po_invoices.replace(db,"special_quote_invoices",quote_id,normalized_invoices,actor_id,actor_name,timestamp)
        else: po_invoices.clear(db,"special_quote_invoices",quote_id,timestamp)
        db.execute("""INSERT INTO special_quote_events(quote_id,event_type,from_status,to_status,follow_up_type,note,user_id,user_name,created_at)
            VALUES(?,'review_saved',?,?,?,?,?,?,?)""",(quote_id,current["status"],status,follow_up_type,"Review saved",actor_id,actor_name,timestamp))
        if changes["status"]: db.execute("""INSERT INTO special_quote_events(quote_id,event_type,from_status,to_status,follow_up_type,note,user_id,user_name,created_at)
            VALUES(?,'status_changed',?,?,?,?,?,?,?)""",(quote_id,current["status"],status,follow_up_type,loss_reason if status=="lost" else "",actor_id,actor_name,timestamp))
        if saved_comment:
            db.execute("INSERT INTO special_quote_comments(quote_id,body,user_id,user_name,created_at) VALUES(?,?,?,?,?)",(quote_id,saved_comment,actor_id,actor_name,timestamp))
            event_type="po_invoices_updated" if automatic and current["status"]=="po" else "po_invoices_registered" if automatic else "comment_added"
            db.execute("INSERT INTO special_quote_events(quote_id,event_type,follow_up_type,note,user_id,user_name,created_at) VALUES(?,?,?,?,?,?,?)",(quote_id,event_type,follow_up_type,saved_comment,actor_id,actor_name,timestamp))
        if current_invoices and status!="po": db.execute("INSERT INTO special_quote_events(quote_id,event_type,follow_up_type,note,user_id,user_name,created_at) VALUES(?,'po_invoices_cleared',?,'Active PO invoices removed after status change',?,?,?)",(quote_id,follow_up_type,actor_id,actor_name,timestamp))
        for changed,event_type,note in ((changes["safe"],"safe_changed","Marked as Safe" if is_safe else "Removed from Safe"),):
            if changed: db.execute("INSERT INTO special_quote_events(quote_id,event_type,follow_up_type,note,user_id,user_name,created_at) VALUES(?,?,?,?,?,?,?)",(quote_id,event_type,follow_up_type,note,actor_id,actor_name,timestamp))
    return get_quote(quote_id)


def restore_quote(quote_id:int)->dict[str,Any]:
    with database.connect() as db:
        db.execute("UPDATE special_quotes SET is_archived=0,archive_reason='Restored by administrator' WHERE id=?",(quote_id,))
    return get_quote(quote_id)


def agents()->list[str]:
    with database.connect() as db: rows=db.execute("SELECT DISTINCT nt_agent FROM special_quotes WHERE trim(nt_agent)<>'' ORDER BY nt_agent").fetchall()
    return [row["nt_agent"] for row in rows]


def ranks()->list[str]:
    with database.connect() as db: rows=db.execute("SELECT DISTINCT rank FROM special_quotes WHERE trim(rank)<>'' ORDER BY rank").fetchall()
    return [row["rank"] for row in rows]


def dashboard(agent:str="",start:str="",end:str="")->dict[str,Any]:
    start,end=database.validate_period(start,end)
    agent_clause=" AND q.nt_agent=?" if agent else ""; agent_params=[agent] if agent else []
    quote_period=" AND q.quote_date BETWEEN ? AND ?" if start else ""
    quote_params=[*agent_params,*([start,end] if start else [])]
    po_period=" AND q.po_date BETWEEN ? AND ?" if start else ""
    po_params=[*agent_params,*([start,end] if start else [])]
    today=database.today_local().isoformat()
    with database.connect() as db:
        pending=dict(db.execute(f"""SELECT COUNT(*) AS all_pending,
            SUM(CASE WHEN q.is_safe=0 THEN 1 ELSE 0 END) AS pending,
            SUM(CASE WHEN q.is_safe=1 THEN 1 ELSE 0 END) AS safe,
            COALESCE(SUM(q.unit_price),0) AS pending_value
            FROM special_quotes q WHERE q.status='pending' AND q.is_archived=0
            {agent_clause}{quote_period}""",quote_params).fetchone())
        po=dict(db.execute(f"""SELECT COUNT(*) AS po,COALESCE(SUM(q.unit_price),0) AS po_quoted_value,
            COALESCE(SUM(COALESCE(q.po_total_usd,0)),0) AS po_value
            FROM special_quotes q WHERE q.status='po' AND q.is_archived=0
            {agent_clause}{po_period}""",po_params).fetchone())
        lost=db.execute(f"SELECT COUNT(*) AS value FROM special_quotes q WHERE q.status='lost' AND q.is_archived=0 {agent_clause}{quote_period}",quote_params).fetchone()["value"]
        pending={key:(value or 0) for key,value in pending.items()}
        po={key:(value or 0) for key,value in po.items()}
        counts={**pending,**po,"lost":lost or 0,"total":int(pending["all_pending"])+int(po["po"])+int(lost or 0)}
        responses_raw=database.rows_to_dicts(db.execute(f"""SELECT q.client_response,COUNT(*) AS count
            FROM special_quotes q WHERE q.status='pending' AND q.is_archived=0 {agent_clause}{quote_period}
            GROUP BY q.client_response""",quote_params).fetchall())
        responses={"yes":0,"no":0,"pending":0}
        for row in responses_raw:
            value=row.get("client_response") or "pending"
            if value in responses: responses[value]+=int(row["count"])
        priorities=database.rows_to_dicts(db.execute(f"""SELECT q.priority,COUNT(*) AS count,COALESCE(SUM(q.unit_price),0) AS value
            FROM special_quotes q WHERE q.status='pending' AND q.is_archived=0 {agent_clause}{quote_period}
            GROUP BY q.priority""",quote_params).fetchall())
        added=db.execute(f"SELECT COUNT(*) AS value FROM special_quotes q WHERE substr(q.discovered_at,1,10)=? AND q.is_archived=0 {agent_clause}",(today,*agent_params)).fetchone()["value"]
        po_today=db.execute(f"""SELECT COUNT(DISTINCT e.quote_id) AS value FROM special_quote_events e
            JOIN special_quotes q ON q.id=e.quote_id WHERE substr(e.created_at,1,10)=? AND e.event_type='status_changed'
            AND e.to_status='po' AND q.is_archived=0 {agent_clause}""",(today,*agent_params)).fetchone()["value"]
        last_import=db.execute("SELECT * FROM imports WHERE workspace='special' AND status='confirmed' ORDER BY id DESC LIMIT 1").fetchone()
        pending_rows=database.rows_to_dicts(db.execute(f"SELECT q.* FROM special_quotes q WHERE q.status='pending' AND q.is_safe=0 AND q.is_archived=0 {agent_clause}",agent_params).fetchall())
        overdue_all=sorted((row for row in (_decorate(row) for row in pending_rows) if row["overdue"]),key=lambda row:row["days_since_activity"],reverse=True)
    return {"currency":"JPY","counts":counts,"responses":responses,"added_today":added,"po_today":po_today,"po_missing":0,"priorities":priorities,
            "last_import":dict(last_import) if last_import else None,"overdue_count":len(overdue_all),"overdue_quotes":overdue_all[:8],
            "start_date":start,"end_date":end}


def report_activity(start_date:str,end_date:str,actor_user_id:int|None=None,agent:str="")->dict[str,Any]:
    start=date.fromisoformat(start_date); end=date.fromisoformat(end_date)
    if end<start: raise ValueError("End date cannot be before start date")
    actor_clause=" AND e.user_id=?" if actor_user_id else ""; agent_clause=" AND q.nt_agent=?" if agent else ""
    event_params=[start_date,end_date]+([actor_user_id] if actor_user_id else [])+([agent] if agent else [])
    quote_params=[start_date,end_date]+([agent] if agent else [])
    created_clause=" AND q.created_by_user_id=?" if actor_user_id else ""; created_params=[start_date,end_date]+([actor_user_id] if actor_user_id else [])+([agent] if agent else [])
    with database.connect() as db:
        events=database.rows_to_dicts(db.execute(f"""SELECT e.*,q.source_quote_number AS folio,q.quote_date,q.customer_name AS end_user,q.nt_agent,q.priority,q.status,
            q.unit_price AS total_usd,q.po_total_usd,q.po_date,q.loss_reason,
            (SELECT COUNT(*) FROM special_quote_invoices i WHERE i.quote_id=q.id AND i.active=1) AS invoice_count
            FROM special_quote_events e JOIN special_quotes q ON q.id=e.quote_id
            WHERE date(e.created_at) BETWEEN ? AND ? {actor_clause} {agent_clause} AND q.is_archived=0 ORDER BY e.created_at,e.id""",event_params).fetchall())
        new_rows=database.rows_to_dicts(db.execute(f"SELECT q.* FROM special_quotes q WHERE date(q.discovered_at) BETWEEN ? AND ? {created_clause} {agent_clause} AND q.is_archived=0",created_params).fetchall())
        pending_rows=database.rows_to_dicts(db.execute(f"SELECT q.* FROM special_quotes q WHERE q.status='pending' AND q.quote_date BETWEEN ? AND ? {agent_clause} AND q.is_archived=0",quote_params).fetchall())
    reviews=[e for e in events if e["event_type"]=="review_saved"]; status_events=[e for e in events if e["event_type"]=="status_changed"]
    lost={e["quote_id"]:e for e in status_events if e["to_status"]=="lost"}; pos={e["quote_id"]:e for e in status_events if e["to_status"]=="po"}
    breakdown:dict[str,int]={}; breakdown_detail:dict[str,dict[str,Any]]={}
    for e in lost.values():
        reason=e.get("note") or e.get("loss_reason") or "Not specified"; breakdown[reason]=breakdown.get(reason,0)+1
        detail=breakdown_detail.setdefault(reason,{"count":0,"value":0.0}); detail["count"]+=1; detail["value"]+=float(e.get("total_usd") or 0)
    reviewed_ids={e["quote_id"] for e in reviews}; reviewed=[]
    for quote_id in reviewed_ids:
        matching=[e for e in events if e["quote_id"]==quote_id and e["event_type"] in {"review_saved","comment_added","status_changed","safe_changed","po_invoices_registered","po_invoices_updated","po_invoices_cleared"}]; base=matching[0]
        reviewed.append({"quote_id":quote_id,"folio":base["folio"] or f"SPQ-{quote_id:05d}","receptor":base["end_user"],"nt_agent":base["nt_agent"],"priority":base["priority"],"status":base["status"],"total":base["total_usd"],"net_total":None,"po_total":base["po_total_usd"],"events":matching})
    po_rows=[{"folio":e["folio"] or f"SPQ-{e['quote_id']:05d}","receptor":e["end_user"],"quote_date":e["quote_date"],"po_date":e["po_date"],"quoted_total":e["total_usd"],"po_total":e["po_total_usd"],"invoice_count":e.get("invoice_count",0),"variance":float(e["po_total_usd"] or 0)-float(e["total_usd"] or 0)} for e in pos.values()]
    lost_rows=[{"folio":e["folio"] or f"SPQ-{e['quote_id']:05d}","receptor":e["end_user"],"quoted_total":e["total_usd"],"reason":e.get("note") or e.get("loss_reason") or "Not specified","date":e["created_at"][:10]} for e in lost.values()]
    pipeline=database.executive_pipeline(pending_rows,"special")
    cycle_days=[]
    for row in po_rows:
        try:
            elapsed=(date.fromisoformat(str(row["po_date"])[:10])-date.fromisoformat(str(row["quote_date"])[:10])).days
            if elapsed>=0: cycle_days.append(elapsed)
        except (TypeError,ValueError):
            pass
    resolved=len(po_rows)+len(lost_rows)
    return {"workspace":"special","currency":"JPY","start_date":start_date,"end_date":end_date,"new_quotes":len(new_rows),
            "quotes_reviewed":len(reviewed_ids),"status_changes":len(status_events),"po_changes":len(po_rows),"lost_changes":len(lost_rows),
            "pending_value":sum(float(row.get("unit_price") or 0) for row in pending_rows),"pending_count":pipeline["pending_count"],
            "new_value":sum(float(row.get("unit_price") or 0) for row in new_rows),
            "lost_value":sum(float(row.get("quoted_total") or 0) for row in lost_rows),
            "conversion_rate":(len(po_rows)/resolved*100) if resolved else None,
            "average_po_days":(sum(cycle_days)/len(cycle_days)) if cycle_days else None,
            "po_quoted_value":sum(float(r["quoted_total"] or 0) for r in po_rows),
            "po_value":sum(float(r["po_total"] or 0) for r in po_rows),"loss_breakdown":breakdown,"po_rows":po_rows,
            "loss_breakdown_detail":breakdown_detail,"lost_rows":lost_rows,"reviewed":reviewed,"new_rows":new_rows,
            **{key:value for key,value in pipeline.items() if key!="pending_count"}}


def monthly_report(months:int=12,agent:str="")->list[dict[str,Any]]:
    clause=" AND nt_agent=?" if agent else ""; params=(agent,) if agent else ()
    with database.connect() as db: rows=db.execute(f"""SELECT substr(quote_date,1,7) AS month,COUNT(*) AS quotes,
        SUM(status='pending') AS pending,SUM(status='po') AS po,SUM(status='lost') AS lost,COALESCE(SUM(unit_price),0) AS quoted_value,
        COALESCE(SUM(CASE WHEN status='po' THEN po_total_usd ELSE 0 END),0) AS po_value FROM special_quotes
        WHERE is_archived=0 {clause} GROUP BY month ORDER BY month DESC LIMIT ?""",(*params,months)).fetchall()
    return list(reversed(database.rows_to_dicts(rows)))


def daily_activity(report_date:str,agent:str="")->dict[str,Any]:
    data=report_activity(report_date,report_date,None,agent)
    data.update({"date":report_date,"agent":agent,"report_currency":"JPY","report_title":"Special quotation daily follow-up report",
                 "activity_count":data["quotes_reviewed"],"quotes":data["reviewed"],"notes":[],"comment_changes":sum(len([e for e in q["events"] if e["event_type"]=="comment_added"]) for q in data["reviewed"]),"reviews_without_changes":0,"follow_up_methods":{m:len({q["quote_id"] for q in data["reviewed"] if any(e["follow_up_type"]==m for e in q["events"])}) for m in ("email","call","visit")}})
    return data
