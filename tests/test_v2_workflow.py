from __future__ import annotations

import sys
import tempfile
import unittest
import base64
import http.client
import json
import threading
from datetime import date, timedelta
from http.server import ThreadingHTTPServer
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from pypdf import PdfReader


PROJECT_DIR=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(PROJECT_DIR))

import auth
import app
import backup
import database
import invoice_manager
import imports_manager
import special
from reports import build_report_pdf, build_report_xlsx


def standard_book(total: float = 6000, order: str = "OP-100") -> bytes:
    book=Workbook(); sheet=book.active; sheet.title="Hoja1"
    sheet.append(["Todos los Documentos"]); sheet.append([])
    sheet.append(["Fecha","Serie","Folio","Razón Social","Total","Texto Extra 1","Texto Extra 3","Pendiente","Texto Extra 2","Nombre del agente","Neto"])
    sheet.append([date.today(),"QT",100,"DISTRIBUTOR SA",total,"END USER MX","DISTRIBUTOR AGENT","",order,"TAKUJI YAMADA",total/1.16])
    sheet.append([date.today(),"QTI",101,"INTERNATIONAL CO",500,"END USER USA","AGENT USA","","","ELEONOR BARRAGAN",431.03])
    buffer=BytesIO(); book.save(buffer); return buffer.getvalue()


def special_book(model_x_price: float = 2000) -> bytes:
    book=Workbook(); sheet=book.active; sheet.title="Sheet1"
    headers=[""]*27
    for index,value in {1:"Quote No.",2:"Quote Date",5:"Company",6:"Rank",7:"Item Code",8:"Model / Size",9:"Qty",10:"Unit Price",27:"Contact"}.items(): headers[index-1]=value
    sheet.append(headers)
    rows=[
        ("S1",date(2026,9,1),"CLIENT A","F","C1","MODEL-X",1,1000,"Ariel"),
        ("S2",date(2026,9,1),"CLIENT A","F","C2","MODEL-X",2,5000,"Ariel"),
        ("S3",date(2026,9,1),"CLIENT A","F","C3","MODEL-X",1,model_x_price,"Ariel"),
        ("S4",date(2026,9,2),"CLIENT A","F","C4","MODEL-X",1,1500,"Ariel"),
        ("S5",date(2026,9,1),"CLIENT B","B","C5","MODEL-Y",5,500,"Eleonor"),
    ]
    for quote_no,quote_date,company,rank,code,model,qty,price,agent in rows:
        values=[""]*27
        for index,value in {1:quote_no,2:quote_date,5:company,6:rank,7:code,8:model,9:qty,10:price,27:agent}.items(): values[index-1]=value
        sheet.append(values)
    buffer=BytesIO(); book.save(buffer); return buffer.getvalue()


def invoice_book(order: str = "OP-100", include_unmatched: bool = True) -> bytes:
    book=Workbook(); sheet=book.active; sheet.title="Facturas"
    sheet.append(["Reporte de facturas"]); sheet.append([])
    sheet.append(["Fecha","Serie","Folio","Cliente","Total","Texto Extra 2"])
    sheet.append([date(2026,9,10),"IV",8772,"CLIENT",150.25,order])
    sheet.append([date(2026,9,12),"IV",8762,"CLIENT",200,order])
    if include_unmatched:
        sheet.append([date(2026,9,13),"IV",9999,"CLIENT",75,"UNKNOWN-PO"])
    buffer=BytesIO(); book.save(buffer); return buffer.getvalue()


class V2WorkflowTests(unittest.TestCase):
    def setUp(self)->None:
        self.temp=tempfile.TemporaryDirectory(); self.old=(database.DATA_DIR,database.DB_PATH,database.IMPORT_DIR,backup.BACKUP_DIR)
        database.DATA_DIR=Path(self.temp.name); database.DB_PATH=database.DATA_DIR/"test.db"; database.IMPORT_DIR=database.DATA_DIR/"imports"; backup.BACKUP_DIR=database.DATA_DIR/"backups"
        database.initialize(); special.initialize(); auth.initialize(); invoice_manager.initialize(); backup.initialize()
        self.user,self.token=auth.bootstrap("takujiyamada","TestPassword123!","en")

    def tearDown(self)->None:
        database.DATA_DIR,database.DB_PATH,database.IMPORT_DIR,backup.BACKUP_DIR=self.old; self.temp.cleanup()

    def test_standard_import_po_missing_comments_and_report(self)->None:
        content=standard_book(); preview=imports_manager.preview("standard",content,"daily.xlsx",self.user)
        self.assertEqual((preview["new"],preview["errors"]),(2,0))
        result=imports_manager.confirm("standard",preview["token"],self.user)
        self.assertEqual((result["inserted"],result["updated"]),(2,0))
        quotes=database.list_quotes({"status":"pending"}); self.assertEqual(len(quotes),2)
        first=next(q for q in quotes if q["folio"]=="QT-100")
        self.assertEqual(first["total_usd"],6000); self.assertAlmostEqual(first["net_total_usd"],6000/1.16)
        self.assertEqual(first["end_user"],"END USER MX"); self.assertEqual(first["priority"],"S"); self.assertEqual(first["po_detected"],1)
        self.assertEqual([q["id"] for q in database.list_quotes({"po_missing":"1"})],[first["id"]])
        with self.assertRaisesRegex(ValueError,"at least one invoice"):
            database.update_quote(first["id"],"po","","","email",False,None,None,self.user,[])
        database.update_quote(first["id"],"pending","Customer contacted","","email",False,None,None,self.user)
        yesterday=(date.today()-timedelta(days=1)).isoformat(); two_days_ago=(date.today()-timedelta(days=2)).isoformat()
        invoices=[
            {"invoice_date":two_days_ago,"invoice_series":"IV","invoice_number":"8,772","amount":3000},
            {"invoice_date":yesterday,"invoice_series":"IV","invoice_number":"8,762","amount":3100},
        ]
        database.update_quote(first["id"],"po","This must not be saved","","call",False,None,None,self.user,invoices)
        detail=database.get_quote(first["id"])
        self.assertEqual(detail["po_total_usd"],6100); self.assertEqual(detail["po_date"],yesterday)
        self.assertEqual(len(detail["invoices"]),2)
        comments=[c["body"] for c in detail["comments"]]
        self.assertEqual(comments[0],"Customer contacted"); self.assertNotIn("This must not be saved",comments)
        self.assertIn("2 invoices",comments[1]); self.assertIn("USD 6,100.00",comments[1])
        self.assertEqual(database.dashboard()["po_today"],1)
        revised=[
            {"invoice_date":two_days_ago,"invoice_series":"IV","invoice_number":"8,772","amount":3000},
            {"invoice_date":yesterday,"invoice_series":"IV","invoice_number":"8,762","amount":3200},
        ]
        database.update_quote(first["id"],"po","","","call",False,None,None,self.user,revised)
        detail=database.get_quote(first["id"]); self.assertEqual(detail["po_total_usd"],6200)
        self.assertEqual(database.dashboard()["po_today"],1)
        report=database.report_activity(date.today().isoformat(),date.today().isoformat(),self.user["id"])
        self.assertEqual(report["quotes_reviewed"],1); self.assertEqual(report["po_changes"],1); self.assertEqual(report["po_value"],6200)
        self.assertEqual(report["po_rows"][0]["invoice_count"],2)
        pdf_path=Path(self.temp.name)/"report.pdf"; pdf=build_report_pdf({**report,"generated_by":self.user["display_name"],"generated_at":database.now_iso(),"scope":"mine","agent":""},pdf_path,"en")
        reader=PdfReader(BytesIO(pdf)); report_text="\n".join(page.extract_text() or "" for page in reader.pages)
        self.assertTrue(pdf.startswith(b"%PDF")); self.assertEqual(len(reader.pages),1)
        self.assertIn("QUOTATION FOLLOW-UP REPORT",report_text); self.assertIn("CONVERSION",report_text); self.assertIn("QUOTATIONS REQUIRING ACTION",report_text)
        self.assertNotIn("Pending pipeline:",report_text)
        xlsx_path=Path(self.temp.name)/"report.xlsx"; build_report_xlsx(report,xlsx_path,"es"); report_book=load_workbook(xlsx_path)
        self.assertEqual(report_book.sheetnames,["Resumen","PO","Perdidas","Revisadas"])
        self.assertEqual(report_book["Resumen"]["A11"].value,"Tasas de conversión")
        self.assertEqual(report_book["Resumen"]["B12"].value,1); self.assertEqual(report_book["Resumen"]["C12"].value,1)
        repeated=imports_manager.preview("standard",content,"daily.xlsx",self.user); self.assertEqual(repeated["updated"],2)

    def test_executive_pdf_handles_empty_results_and_overflow_actions(self)->None:
        action_rows=[{
            "priority":"S" if index<3 else "A","folio":f"QT-{index:04d}","distributor_code":"ABC",
            "end_user":"CUSTOMER WITH A VERY LONG CORPORATE NAME THAT MUST BE TRUNCATED SAFELY",
            "amount":99999999+index,"days_since_activity":91+index,"nt_agent":"AGENT WITH A LONG DISPLAY NAME",
        } for index in range(12)]
        base={
            "start_date":"2026-01-01","end_date":"2026-09-24","scope":"all","agent":"",
            "generated_by":"TEST USER","generated_at":database.now_iso(),"pending_count":12,
            "pending_value":1199999994,"new_quotes":0,"new_value":0,"quotes_reviewed":0,
            "po_changes":0,"po_value":0,"lost_changes":0,"lost_value":0,"conversion_rate":None,
            "po_quoted_value":0,"average_po_days":None,"priorities":[],"loss_breakdown_detail":{},
            "alerts":{"followup_9_14":{"count":0,"value":0},"followup_15_plus":{"count":12,"value":1199999994},
                      "stale_90_plus":{"count":12,"value":1199999994},"po_date_missing":{"count":0,"value":0}},
            "action_rows":action_rows,"action_remaining":2,
        }
        for workspace,currency,language in (("standard","USD","en"),("special","JPY","es")):
            pdf=build_report_pdf({**base,"workspace":workspace,"currency":currency},Path(self.temp.name)/f"{workspace}-overflow.pdf",language)
            reader=PdfReader(BytesIO(pdf)); text="\n".join(page.extract_text() or "" for page in reader.pages)
            self.assertEqual(len(reader.pages),1); self.assertIn("N/A",text); self.assertIn("+2",text)

    def test_executive_pipeline_alert_boundaries_and_deduplication(self)->None:
        today=date.today()
        rows=[]
        for index,days in enumerate((8,9,14,15,90,91),start=1):
            rows.append({
                "id":index,"folio":f"QT-{index}","quote_date":(today-timedelta(days=days)).isoformat(),
                "discovered_at":(today-timedelta(days=days)).isoformat(),"last_reviewed_at":None,
                "updated_at":(today-timedelta(days=days)).isoformat(),"status":"pending","is_safe":0,
                "priority":"S","total_usd":1000*index,"distributor_code":"ABC","end_user":"CLIENT",
                "nt_agent":"AGENT","po_detected":1 if days==91 else 0,"po_date":None,
            })
        pipeline=database.executive_pipeline(rows,"standard")
        self.assertEqual(pipeline["alerts"]["followup_9_14"]["count"],2)
        self.assertEqual(pipeline["alerts"]["followup_15_plus"]["count"],3)
        self.assertEqual(pipeline["alerts"]["stale_90_plus"]["count"],1)
        self.assertEqual(pipeline["alerts"]["po_date_missing"]["count"],1)
        self.assertEqual(len(pipeline["action_rows"]),4)

    def test_special_duplicate_visibility_and_archive(self)->None:
        content=special_book(); records,errors=special.parse_workbook(content,"special.xlsx")
        self.assertEqual(errors,[]); visible=[r for r in records if not r["is_archived"]]; archived=[r for r in records if r["is_archived"]]
        self.assertEqual(len(visible),3); self.assertEqual(len(archived),2)
        self.assertIn("S3",[r["source_quote_number"] for r in visible]); self.assertIn("S4",[r["source_quote_number"] for r in visible]); self.assertIn("S5",[r["source_quote_number"] for r in visible])
        preview=imports_manager.preview("special",content,"special.xlsx",self.user); self.assertEqual(preview["archived"],2)
        first_import=imports_manager.confirm("special",preview["token"],self.user)
        self.assertEqual((first_import["inserted"],first_import["updated"],first_import["unchanged"]),(5,0,0))
        self.assertEqual(len(special.list_quotes({})),3); self.assertEqual(len(special.list_quotes({"archive":"1"})),2)
        quote=special.list_quotes({})[0]
        special.update_quote(quote["id"],"lost","Budget rejected","Over Budget","call",False,None,None,self.user)
        self.assertEqual(special.get_quote(quote["id"])["comments"][0]["body"],"Budget rejected")
        report=special.report_activity(date.today().isoformat(),date.today().isoformat(),self.user["id"])
        self.assertEqual(report["currency"],"JPY"); self.assertEqual(report["lost_changes"],1); self.assertEqual(report["loss_breakdown"],{"Over Budget":1})
        repeated=imports_manager.preview("special",content,"special.xlsx",self.user)
        self.assertEqual((repeated["new"],repeated["updated"],repeated["unchanged"]),(0,0,5))
        repeated_result=imports_manager.confirm("special",repeated["token"],self.user)
        self.assertEqual((repeated_result["inserted"],repeated_result["updated"],repeated_result["unchanged"]),(0,0,5))
        changed_content=special_book(2250)
        changed=imports_manager.preview("special",changed_content,"special_changed.xlsx",self.user)
        self.assertEqual((changed["new"],changed["updated"],changed["unchanged"]),(0,1,4))
        changed_result=imports_manager.confirm("special",changed["token"],self.user)
        self.assertEqual((changed_result["inserted"],changed_result["updated"],changed_result["unchanged"]),(0,1,4))
        changed_quote=next(row for row in special.list_quotes({}) if row["source_quote_number"]=="S3")
        self.assertEqual(changed_quote["unit_price"],2250)

    def test_managed_and_unmanaged_views_are_consistent_in_both_workspaces(self)->None:
        standard_preview=imports_manager.preview("standard",standard_book(),"daily.xlsx",self.user)
        imports_manager.confirm("standard",standard_preview["token"],self.user)
        standard_rows=database.list_quotes({})
        self.assertEqual(database.list_managed_quotes({}),[])
        self.assertEqual(len(database.list_management_quotes({"status":"unmanaged"})),2)
        standard_quote=next(row for row in standard_rows if row["folio"]=="QTI-101")
        database.update_quote(standard_quote["id"],"pending","","","email",False,None,None,self.user)
        self.assertEqual([row["id"] for row in database.list_managed_quotes({})],[standard_quote["id"]])
        self.assertNotIn(standard_quote["id"],[row["id"] for row in database.list_management_quotes({"status":"unmanaged"})])

        special_preview=imports_manager.preview("special",special_book(),"special.xlsx",self.user)
        imports_manager.confirm("special",special_preview["token"],self.user)
        special_rows=special.list_quotes({})
        self.assertEqual(special.list_managed({}),[])
        self.assertEqual(len(special.list_management({"status":"unmanaged"})),3)
        special_quote=next(row for row in special_rows if row["source_quote_number"]=="S3")
        special.update_quote(special_quote["id"],"pending","","","visit",False,None,None,self.user)
        self.assertEqual([row["id"] for row in special.list_managed({})],[special_quote["id"]])
        self.assertNotIn(special_quote["id"],[row["id"] for row in special.list_management({"status":"unmanaged"})])

        with database.connect() as db:
            db.execute("UPDATE quotes SET last_reviewed_at=NULL,last_reviewed_by=NULL WHERE id=?",(standard_quote["id"],))
            db.execute("UPDATE special_quotes SET last_reviewed_at=NULL,last_reviewed_by=NULL WHERE id=?",(special_quote["id"],))
        database.initialize(); special.initialize()
        self.assertEqual([row["id"] for row in database.list_managed_quotes({})],[standard_quote["id"]])
        self.assertEqual([row["id"] for row in special.list_managed({})],[special_quote["id"]])

    def test_special_po_uses_multiple_jpy_invoices(self)->None:
        preview=imports_manager.preview("special",special_book(),"special.xlsx",self.user)
        imports_manager.confirm("special",preview["token"],self.user)
        quote=next(row for row in special.list_quotes({}) if row["source_quote_number"]=="S3")
        invoice_date=(date.today()-timedelta(days=3)).isoformat()
        result=special.update_quote(quote["id"],"po","Ignored","","visit",False,None,None,self.user,[
            {"invoice_date":invoice_date,"invoice_series":"IV","invoice_number":"1001","amount":800000},
            {"invoice_date":invoice_date,"invoice_series":"IV","invoice_number":"1002","amount":250000},
        ])
        self.assertEqual(result["po_total_usd"],1050000); self.assertEqual(result["po_date"],invoice_date)
        self.assertEqual(len(result["invoices"]),2); self.assertIn("JPY 1,050,000",result["comments"][0]["body"])
        self.assertEqual(special.dashboard()["po_today"],1)
        report=special.report_activity(date.today().isoformat(),date.today().isoformat(),self.user["id"])
        self.assertEqual(report["po_value"],1050000); self.assertEqual(report["po_rows"][0]["invoice_count"],2)
        pdf=build_report_pdf({**report,"generated_by":self.user["display_name"],"generated_at":database.now_iso(),"scope":"mine","agent":""},Path(self.temp.name)/"special-report.pdf","es")
        reader=PdfReader(BytesIO(pdf)); report_text="\n".join(page.extract_text() or "" for page in reader.pages)
        self.assertEqual(len(reader.pages),1); self.assertIn("JPY",report_text); self.assertIn("OPERACIONES QUE REQUIEREN",report_text)

    def test_historical_po_totals_migrate_to_legacy_invoices(self)->None:
        standard_preview=imports_manager.preview("standard",standard_book(),"daily.xlsx",self.user)
        imports_manager.confirm("standard",standard_preview["token"],self.user)
        standard_quote=next(q for q in database.list_quotes({}) if q["folio"]=="QT-100")
        special_preview=imports_manager.preview("special",special_book(),"special.xlsx",self.user)
        imports_manager.confirm("special",special_preview["token"],self.user)
        special_quote=next(q for q in special.list_quotes({}) if q["source_quote_number"]=="S3")
        legacy_date=(date.today()-timedelta(days=30)).isoformat()
        with database.connect() as db:
            db.execute("UPDATE quotes SET status='po',po_total_usd=4321,po_date=? WHERE id=?",(legacy_date,standard_quote["id"]))
            db.execute("UPDATE special_quotes SET status='po',po_total_usd=765432,po_date=? WHERE id=?",(legacy_date,special_quote["id"]))
        database.initialize(); special.initialize()
        standard_detail=database.get_quote(standard_quote["id"]); special_detail=special.get_quote(special_quote["id"])
        self.assertEqual(len(standard_detail["invoices"]),1); self.assertEqual(standard_detail["invoices"][0]["is_legacy"],1)
        self.assertEqual(standard_detail["invoices"][0]["amount"],4321)
        self.assertEqual(len(special_detail["invoices"]),1); self.assertEqual(special_detail["invoices"][0]["amount"],765432)
        database.initialize(); special.initialize()
        self.assertEqual(len(database.get_quote(standard_quote["id"])["invoices"]),1)
        self.assertEqual(len(special.get_quote(special_quote["id"])["invoices"]),1)

    def test_invoice_excel_replaces_legacy_invoice_and_is_idempotent(self)->None:
        quotation=imports_manager.preview("standard",standard_book(),"daily.xlsx",self.user)
        imports_manager.confirm("standard",quotation["token"],self.user)
        quote=next(row for row in database.list_quotes({}) if row["folio"]=="QT-100")
        legacy_date=date(2026,8,1).isoformat()
        with database.connect() as db:
            db.execute("UPDATE quotes SET status='po',follow_up_type='email',po_total_usd=9999,po_date=? WHERE id=?",(legacy_date,quote["id"]))
        database.initialize()
        self.assertEqual(database.get_quote(quote["id"])["invoices"][0]["is_legacy"],1)

        preview=invoice_manager.preview("standard",invoice_book(),"invoices.xlsx",self.user)
        self.assertEqual((preview["new"],preview["unmatched"],preview["invalid"]),(2,1,0))
        result=invoice_manager.confirm("standard",preview["token"],self.user)
        self.assertEqual((result["invoices_added"],result["quotes_updated"],result["legacy_invoices_replaced"]),(2,1,1))
        detail=database.get_quote(quote["id"])
        self.assertEqual([row["invoice_number"] for row in detail["invoices"]],["8772","8762"])
        self.assertTrue(all(not row["is_legacy"] for row in detail["invoices"]))
        self.assertAlmostEqual(detail["po_total_usd"],350.25)
        self.assertEqual(detail["po_date"],"2026-09-12")
        self.assertIn("USD 350.25",detail["comments"][-1]["body"])
        self.assertIn("po_invoices_registered",[event["event_type"] for event in detail["events"]])
        self.assertEqual(detail["events"][-1]["event_type"],"status_changed")

        repeated=invoice_manager.preview("standard",invoice_book(include_unmatched=False),"invoices-again.xlsx",self.user)
        self.assertEqual((repeated["new"],repeated["duplicate"]),(0,2))
        repeated_result=invoice_manager.confirm("standard",repeated["token"],self.user)
        self.assertEqual((repeated_result["invoices_added"],repeated_result["quotes_updated"]),(0,0))
        with self.assertRaisesRegex(ValueError,"already been confirmed"):
            invoice_manager.confirm("standard",repeated["token"],self.user)

    def test_invoice_excel_accepts_detected_orders_and_rejects_special_workspace(self)->None:
        quotation=imports_manager.preview("standard",standard_book(),"daily.xlsx",self.user)
        imports_manager.confirm("standard",quotation["token"],self.user)
        preview=invoice_manager.preview("standard",invoice_book(include_unmatched=False),"invoices.xlsx",self.user)
        self.assertEqual((preview["new"],preview["not_po"]),(2,0))
        with self.assertRaisesRegex(ValueError,"only in Follow Up Quotations"):
            invoice_manager.preview("special",invoice_book(),"invoices.xlsx",self.user)

    def test_invoice_import_http_route_accepts_authenticated_csrf_request(self)->None:
        server=ThreadingHTTPServer(("127.0.0.1",0),app.Handler)
        thread=threading.Thread(target=server.serve_forever,daemon=True); thread.start()
        connection=http.client.HTTPConnection("127.0.0.1",server.server_port,timeout=10)
        headers={
            "Content-Type":"application/json",
            "Cookie":f"{app.COOKIE_NAME}={self.token}",
            "X-CSRF-Token":self.user["csrf_token"],
        }
        try:
            body=json.dumps({"filename":"invoices.xlsx","content_base64":base64.b64encode(invoice_book()).decode("ascii")})
            connection.request("POST","/api/import/invoices/preview",body=body,headers=headers)
            response=connection.getresponse(); payload=json.loads(response.read())
            self.assertEqual(response.status,201); self.assertEqual(payload["unmatched"],3)
            connection.request("POST","/api/import/invoices/confirm",body=json.dumps({"token":payload["token"]}),headers=headers)
            response=connection.getresponse(); confirmed=json.loads(response.read())
            self.assertEqual(response.status,201); self.assertEqual(confirmed["invoices_added"],0)
        finally:
            connection.close(); server.shutdown(); server.server_close(); thread.join(timeout=5)

    def test_special_legacy_duplicate_is_merged_with_history(self)->None:
        content=special_book(); preview=imports_manager.preview("special",content,"special.xlsx",self.user)
        imports_manager.confirm("special",preview["token"],self.user)
        with database.connect() as db:
            current=db.execute("SELECT * FROM special_quotes WHERE source_quote_number='S3'").fetchone()
            values=dict(current); values.pop("id"); values.update({
                "source_key":"legacy-full-row-key","source_quote_number":"","customer_name":"",
                "source_file":"old-special.xlsx","import_id":None,"status":"lost","loss_reason":"Over Budget",
                "follow_up_type":"call","comment":"Legacy follow-up","last_reviewed_at":database.now_iso(),
            })
            columns=list(values); marks=",".join("?" for _ in columns)
            legacy_id=db.execute(f"INSERT INTO special_quotes({','.join(columns)}) VALUES({marks})",tuple(values[c] for c in columns)).lastrowid
            db.execute("INSERT INTO special_quote_comments(quote_id,body,user_name,created_at) VALUES(?,?,?,?)",
                       (legacy_id,"Legacy follow-up","Historical data",database.now_iso()))
        special.initialize()
        with database.connect() as db:
            legacy=db.execute("SELECT * FROM special_quotes WHERE id=?",(legacy_id,)).fetchone()
            active=db.execute("SELECT * FROM special_quotes WHERE source_quote_number='S3' AND is_archived=0").fetchall()
        self.assertEqual(len(active),1); self.assertEqual(active[0]["status"],"lost")
        self.assertEqual(legacy["is_archived"],1); self.assertIn("Merged duplicate",legacy["archive_reason"])
        detail=special.get_quote(active[0]["id"])
        self.assertIn("Legacy follow-up",[comment["body"] for comment in detail["comments"]])

    def test_auth_user_password_reset_backup_and_audit(self)->None:
        logged,token=auth.login("takujiyamada","TestPassword123!","es")
        self.assertEqual(logged["language"],"es"); self.assertIsNotNone(auth.authenticate(token))
        created=auth.create_user(logged,{"username":"viewer","display_name":"Viewer","role":"superadmin","password":"ViewerPass123!","agent_name":"Viewer"})
        self.assertEqual(created["role"],"user"); self.assertTrue(created["can_edit"])
        updated=auth.update_user(logged,created["id"],{"display_name":"Viewer","role":"superadmin","agent_name":"Viewer","active":True,"password":"TemporaryPass456!"})
        self.assertEqual(updated["role"],"user"); self.assertTrue(updated["must_change_password"])
        reset_login,_=auth.login("viewer","TemporaryPass456!","en"); self.assertTrue(reset_login["must_change_password"])
        result=backup.perform(logged["display_name"]); self.assertTrue(result["ok"]); self.assertTrue(result["verified"]); self.assertFalse(backup.status()["overdue"])
        self.assertFalse(logged["can_view_audit"])
        with self.assertRaisesRegex(PermissionError,"Secretadmin"):
            auth.recent_audit(logged)
        ariel=next(user for user in auth.list_users(logged) if user["username"]=="arielcontreras")
        auth.update_user(logged,ariel["id"],{"display_name":ariel["display_name"],"agent_name":ariel["agent_name"],"active":True,"password":"SecretAdmin123!"})
        secret,_=auth.login("arielcontreras","SecretAdmin123!","en")
        self.assertTrue(secret["can_view_audit"])
        audit=auth.recent_audit(secret); self.assertTrue(any(row["action"]=="user_created" for row in audit)); self.assertTrue(any(row["action"]=="user_updated" for row in audit))

    def test_language_switch_preserves_manage_and_user_controls(self)->None:
        html=(PROJECT_DIR/"static"/"index.html").read_text(encoding="utf-8")
        javascript=(PROJECT_DIR/"static"/"app.js").read_text(encoding="utf-8")
        for element_id in ("sidebar-toggle","manage-status","manage-method","manage-comment","add-invoice","invoice-rows","new-display-name","new-agent","new-password","edit-name","edit-agent","edit-password"):
            self.assertIn(f'id="{element_id}"',html)
        self.assertIn("invoicePayload()",javascript)
        self.assertIn("/api/import/invoices/preview",javascript)
        self.assertNotIn("await loadData()",javascript)
        self.assertNotIn("const request = await fetch(endpoint",javascript)
        self.assertIn("sidebar-collapsed",javascript)
        self.assertIn("el.matches('label')",javascript)
        self.assertIn(":scope > input, :scope > select, :scope > textarea",javascript)
        self.assertNotIn("forEach(el=>el.textContent=t(el.dataset.i18n))",javascript)

    def test_calendar_boundaries_and_frontend_regressions(self)->None:
        self.assertEqual(database.calendar_periods(date(2026,3,31))["fiscal_start"],"2025-04-01")
        self.assertEqual(database.calendar_periods(date(2026,4,1))["fiscal_start"],"2026-04-01")
        html=(PROJECT_DIR/"static"/"index.html").read_text(encoding="utf-8")
        javascript=(PROJECT_DIR/"static"/"app.js").read_text(encoding="utf-8")
        self.assertEqual(javascript.count("async function selectSystem"),1)
        self.assertEqual(javascript.count("$('#user-edit-form').addEventListener('submit'"),1)
        self.assertNotIn("confirm(",javascript)
        self.assertIn('id="stale-warning"',html)
        self.assertIn('<option value="ja">',html)
        self.assertIn("Perfil de visualización",javascript)
        japanese=auth.update_preferences(self.user,"ja")
        self.assertEqual(japanese["language"],"ja")

    def test_client_response_is_persisted_in_both_workspaces(self)->None:
        standard_preview=imports_manager.preview("standard",standard_book(),"daily.xlsx",self.user)
        imports_manager.confirm("standard",standard_preview["token"],self.user)
        standard_quote=next(row for row in database.list_quotes({}) if row["folio"]=="QTI-101")
        updated=database.update_quote(standard_quote["id"],"pending","","","email",False,None,None,self.user,client_response="yes")
        self.assertEqual(updated["client_response"],"yes")
        self.assertEqual(database.dashboard()["responses"]["yes"],1)
        with self.assertRaisesRegex(ValueError,"Invalid client response"):
            database.update_quote(standard_quote["id"],"pending","","","email",False,None,None,self.user,client_response="maybe")

        special_preview=imports_manager.preview("special",special_book(),"special.xlsx",self.user)
        imports_manager.confirm("special",special_preview["token"],self.user)
        special_quote=next(row for row in special.list_quotes({}) if row["source_quote_number"]=="S3")
        updated=special.update_quote(special_quote["id"],"pending","","","call",False,None,None,self.user,client_response="no")
        self.assertEqual(updated["client_response"],"no")
        self.assertEqual(special.dashboard()["responses"]["no"],1)

    def test_managed_status_grouping_and_activity_timestamp(self)->None:
        preview=imports_manager.preview("standard",standard_book(),"daily.xlsx",self.user)
        imports_manager.confirm("standard",preview["token"],self.user)
        rows=database.list_quotes({})
        pending=next(row for row in rows if row["folio"]=="QTI-101")
        converted=next(row for row in rows if row["folio"]=="QT-100")
        database.update_quote(pending["id"],"pending","","","email",False,None,None,self.user)
        database.update_quote(converted["id"],"lost","","Over Budget","call",False,None,None,self.user)
        managed=database.list_managed_quotes({})
        self.assertEqual([row["status"] for row in managed],["pending","lost"])
        lost=next(row for row in managed if row["status"]=="lost")
        self.assertEqual(lost["last_activity_at"],lost["status_changed_at"])
        self.assertEqual(lost["quote_age_days"],database._days_since(lost["quote_date"]))

    def test_special_managed_returns_every_reviewed_status_group(self)->None:
        preview=imports_manager.preview("special",special_book(),"special.xlsx",self.user)
        imports_manager.confirm("special",preview["token"],self.user)
        rows=special.list_quotes({})
        self.assertEqual(len(rows),3)
        special.update_quote(rows[0]["id"],"pending","","","email",False,None,None,self.user)
        special.update_quote(rows[1]["id"],"lost","","Mismatch","call",False,None,None,self.user)
        special.update_quote(rows[2]["id"],"po","","","visit",False,None,None,self.user,[{
            "invoice_date":date.today().isoformat(),"invoice_series":"IV","invoice_number":"JP-1","amount":500,
        }])
        managed=special.list_managed({})
        self.assertEqual(len(managed),3)
        self.assertEqual([row["status"] for row in managed],["pending","lost","po"])
        self.assertEqual(special.list_management({"status":"unmanaged"}),[])

    def test_dashboard_period_uses_quote_date_for_pending_and_po_date_for_po(self)->None:
        preview=imports_manager.preview("standard",standard_book(),"daily.xlsx",self.user)
        imports_manager.confirm("standard",preview["token"],self.user)
        quote=next(row for row in database.list_quotes({}) if row["folio"]=="QT-100")
        po_day=(date.today()-timedelta(days=1)).isoformat()
        database.update_quote(quote["id"],"po","","","email",False,None,None,self.user,[{
            "invoice_date":po_day,"invoice_series":"IV","invoice_number":"USD-1","amount":6100,
        }])
        previous=database.dashboard(start=po_day,end=po_day)
        current=database.dashboard(start=date.today().isoformat(),end=date.today().isoformat())
        self.assertEqual(previous["counts"]["po"],1)
        self.assertEqual(previous["counts"]["pending"],0)
        self.assertEqual(current["counts"]["po"],0)
        self.assertEqual(current["counts"]["pending"],1)

    def test_report_period_activity_uses_current_pipeline_and_cumulative_losses(self)->None:
        preview=imports_manager.preview("standard",standard_book(),"daily.xlsx",self.user)
        imports_manager.confirm("standard",preview["token"],self.user)
        rows=database.list_quotes({})
        lost=next(row for row in rows if row["folio"]=="QT-100")
        pending=next(row for row in rows if row["folio"]=="QTI-101")
        database.update_quote(lost["id"],"lost","","Over Budget","call",False,None,None,self.user)
        old_date=(date.today()-timedelta(days=40)).isoformat()
        with database.connect() as db:
            db.execute("UPDATE quotes SET quote_date=?,discovered_at=? WHERE id=?",(old_date,old_date,pending["id"]))
        future=(date.today()+timedelta(days=1)).isoformat()
        report=database.report_activity(future,future,None,"")
        self.assertEqual(report["new_quotes"],0)
        self.assertEqual(report["lost_changes"],0)
        self.assertEqual(report["pending_count"],1)
        self.assertEqual(report["pending_value"],500)
        self.assertEqual(report["loss_breakdown"],{"Over Budget":1})

    def test_unmatched_distributor_is_logged_and_details_keep_code(self)->None:
        preview=imports_manager.preview("standard",standard_book(),"daily.xlsx",self.user)
        imports_manager.confirm("standard",preview["token"],self.user)
        with database.connect() as db:
            unmatched=db.execute("SELECT company_name,occurrences FROM distributor_unmatched WHERE normalized_name=?",("distributor sa",)).fetchone()
        self.assertIsNotNone(unmatched)
        self.assertEqual(unmatched["company_name"],"DISTRIBUTOR SA")
        javascript=(PROJECT_DIR/"static"/"app.js").read_text(encoding="utf-8")
        self.assertIn("q.distributor_code||t('no_code')",javascript)


if __name__=="__main__": unittest.main()
