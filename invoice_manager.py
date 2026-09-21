from __future__ import annotations

import io
import json
import math
import re
import secrets
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

import auth
import database
import po_invoices


REQUIRED_HEADERS = {"fecha", "serie", "folio", "texto extra 2", "total"}


def initialize() -> None:
    """Create persistent storage for invoice-import previews."""
    with database.connect() as db:
        db.executescript(
            """
            CREATE TABLE IF NOT EXISTS invoice_import_batches (
                token TEXT PRIMARY KEY,
                workspace TEXT NOT NULL,
                filename TEXT NOT NULL,
                payload_json TEXT NOT NULL,
                preview_json TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'preview',
                user_id INTEGER NOT NULL,
                user_name TEXT NOT NULL,
                created_at TEXT NOT NULL,
                confirmed_at TEXT
            );
            CREATE INDEX IF NOT EXISTS idx_invoice_import_batches_user
            ON invoice_import_batches(user_id,status,created_at);
            """
        )
        #--- NUEVO: Sincronizar el gafete de Excel retroactivamente---
        batches =db.execute("SELECT payload_json FROM invoice_import_batches WHERE status='confirmed'").fetchall()
        for b in batches:
            try:
                records = json.loads(b["payload_json"])
                for r in records:
                    qid = int(r["quote_id"])
                    #Solo lo inserta si no existe ya
                    db.execute("INSERT INTO po_invoices (quote_id) SELECT ? WHERE NOT EXISTS(SELECT 1 FROM po_invoices WHERE quote_id=?)", (qid, qid))
            except Exception:
                pass


def _normalize(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    return " ".join(text.strip().casefold().split())


def _identifier(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _amount(value: Any, row_number: int) -> float:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        amount = float(value)
    else:
        text = str(value or "").strip().replace("$", "").replace("¥", "").replace(" ", "")
        if not text:
            raise ValueError(f"Row {row_number}: Total is empty")
        if "," in text and "." in text:
            text = text.replace(",", "") if text.rfind(".") > text.rfind(",") else text.replace(".", "").replace(",", ".")
        elif "," in text:
            tail = text.rsplit(",", 1)[-1]
            text = text.replace(",", ".") if len(tail) in {1, 2} else text.replace(",", "")
        amount = float(text)
    if not math.isfinite(amount) or amount <= 0:
        raise ValueError(f"Row {row_number}: Total must be greater than zero")
    return amount


def _date(value: Any, row_number: int) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    if not text:
        raise ValueError(f"Row {row_number}: Fecha is empty")
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"Row {row_number}: Fecha is invalid")


def _workbook_rows(content: bytes, filename: str) -> list[list[Any]]:
    suffix = Path(filename).suffix.casefold()
    if suffix == ".xls":
        try:
            import xlrd  # type: ignore
        except ImportError as exc:
            raise ValueError("Legacy .xls support is unavailable. Save the file as .xlsx and try again.") from exc
        try:
            book = xlrd.open_workbook(file_contents=content)
            sheet = book.sheet_by_index(0)
            rows: list[list[Any]] = []
            for row_index in range(sheet.nrows):
                values: list[Any] = []
                for col_index in range(sheet.ncols):
                    cell = sheet.cell(row_index, col_index)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        values.append(datetime(*xlrd.xldate_as_tuple(cell.value, book.datemode)))
                    else:
                        values.append(cell.value)
                rows.append(values)
            return rows
        except Exception as exc:
            raise ValueError(f"The .xls file could not be read: {exc}") from exc
    if suffix not in {".xlsx", ".xlsm"}:
        raise ValueError("Upload an Excel file in .xlsx, .xlsm, or .xls format")
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = book.active
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    except Exception as exc:
        raise ValueError(f"The Excel file could not be read: {exc}") from exc


def _parse(content: bytes, filename: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    rows = _workbook_rows(content, filename)
    header_index = -1
    positions: dict[str, int] = {}
    for index, row in enumerate(rows[:25]):
        candidate = {_normalize(value): column for column, value in enumerate(row) if _normalize(value)}
        if REQUIRED_HEADERS.issubset(candidate):
            header_index = index
            positions = candidate
            break
    if header_index < 0:
        names = ", ".join(sorted(REQUIRED_HEADERS))
        raise ValueError(f"The Excel file must contain these headers: {names}")

    parsed: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for source_index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        def cell(name: str) -> Any:
            column = positions[name]
            return row[column] if column < len(row) else None

        customer_order = _identifier(cell("texto extra 2"))
        if not customer_order:
            continue
        try:
            invoice_series = _identifier(cell("serie"))
            invoice_number = _identifier(cell("folio"))
            if not invoice_series or not invoice_number:
                raise ValueError(f"Row {source_index}: Serie and Folio are required")
            parsed.append(
                {
                    "source_row": source_index,
                    "customer_order": customer_order,
                    "invoice_date": _date(cell("fecha"), source_index),
                    "invoice_series": invoice_series,
                    "invoice_number": invoice_number,
                    "amount": _amount(cell("total"), source_index),
                }
            )
        except (TypeError, ValueError) as exc:
            errors.append({"row": source_index, "error": str(exc)})
    if not parsed and not errors:
        raise ValueError("The Excel file does not contain invoice rows with Texto Extra 2")
    return parsed, errors


def _safe_filename(filename: str) -> str:
    return re.sub(r"[^A-Za-z0-9._ -]+", "_", Path(filename).name)[:160] or "invoices.xlsx"


def preview(workspace: str, content: bytes, filename: str, user: dict[str, Any]) -> dict[str, Any]:
    if workspace != "standard":
        raise ValueError("Invoice Excel import is available only in Follow Up Quotations (USD)")
    if user.get("role") == "readonly":
        raise PermissionError("Read-only users cannot import invoices")
    if not content:
        raise ValueError("Select an Excel file")
    initialize()
    filename = _safe_filename(filename)
    parsed, errors = _parse(content, filename)

    with database.connect() as db:
        quotes = db.execute(
            """SELECT id,folio,status,customer_order,po_detected,po_date FROM quotes
            WHERE is_historical=0 AND is_archived=0 AND TRIM(customer_order)<>''"""
        ).fetchall()
        matches: dict[str, list[dict[str, Any]]] = {}
        for quote in quotes:
            matches.setdefault(_normalize(quote["customer_order"]), []).append(dict(quote))

        active_rows = db.execute(
            """SELECT quote_id,invoice_series,invoice_number FROM quote_invoices WHERE active=1"""
        ).fetchall()
        existing: dict[int, set[tuple[str, str]]] = {}
        for invoice in active_rows:
            existing.setdefault(int(invoice["quote_id"]), set()).add(
                (_normalize(invoice["invoice_series"]), _normalize(invoice["invoice_number"]))
            )

    new_records: list[dict[str, Any]] = []
    result_rows: list[dict[str, Any]] = []
    upload_seen: dict[tuple[int, str, str], dict[str, Any]] = {}
    counts = {"new": 0, "duplicate": 0, "unmatched": 0, "ambiguous": 0, "not_po": 0, "invalid": len(errors)}
    
    for row in parsed:
        candidates = matches.get(_normalize(row["customer_order"]), [])
        # CAMBIO 1: Aceptamos CUALQUIER cotización Pendiente (tenga o no po_detected) que no tenga fecha PO
        po_candidates = [
            quote for quote in candidates
            if quote["status"] == "pending" and not quote["po_date"]
        ]
        result = {**row, "quote_folio": "", "status": ""}
        
        if not candidates:
            result["status"] = "unmatched"
            counts["unmatched"] += 1
        elif not po_candidates:
            result["status"] = "not_po"
            result["quote_folio"] = candidates[0]["folio"]
            counts["not_po"] += 1
        else:
            # CAMBIO 2: Si hay varias cotizaciones con la misma PO, inyectamos la factura a TODAS
            added_to_any = False
            is_duplicate = False
            
            for quote in po_candidates:
                quote_id = int(quote["id"])
                key = (quote_id, _normalize(row["invoice_series"]), _normalize(row["invoice_number"]))
                
                if key in upload_seen:
                    prior = upload_seen[key]
                    if prior["invoice_date"] == row["invoice_date"] and abs(float(prior["amount"]) - float(row["amount"])) < 0.005:
                        is_duplicate = True
                elif key[1:] in existing.get(quote_id, set()):
                    is_duplicate = True
                else:
                    upload_seen[key] = row
                    new_records.append({**row, "quote_id": quote_id, "quote_folio": quote["folio"]})
                    added_to_any = True
            
            if added_to_any:
                result["status"] = "new"
                result["quote_folio"] = po_candidates[0]["folio"] + (" (Varias)" if len(po_candidates) > 1 else "")
                counts["new"] += 1
            else:
                result["status"] = "duplicate"
                counts["duplicate"] += 1
                result["quote_folio"] = po_candidates[0]["folio"]
        result_rows.append(result)

    token = secrets.token_urlsafe(24)
    preview_data = {**counts, "rows_seen": len(parsed) + len(errors), "rows": result_rows, "errors": errors}
    with database.connect() as db:
        db.execute(
            """INSERT INTO invoice_import_batches(token,workspace,filename,payload_json,preview_json,status,
            user_id,user_name,created_at) VALUES(?,?,?,?,?,'preview',?,?,?)""",
            (token, workspace, filename, json.dumps(new_records, ensure_ascii=False),
             json.dumps(preview_data, ensure_ascii=False), user["id"], user["display_name"], database.now_iso()),
        )
    auth.audit(user, "invoice_import_previewed", workspace, "invoice_import", token,
               f"{filename}: {counts['new']} new, {counts['duplicate']} duplicate, {counts['unmatched']} unmatched")
    return {"token": token, "filename": filename, **preview_data}

def confirm(workspace: str, token: str, user: dict[str, Any]) -> dict[str, Any]:
    if workspace != "standard":
        raise ValueError("Invoice Excel import is available only in Follow Up Quotations (USD)")
    if user.get("role") == "readonly":
        raise PermissionError("Read-only users cannot import invoices")
    initialize()
    timestamp = database.now_iso()
    invoices_added = 0
    quotes_updated = 0
    legacy_replaced = 0

    with database.connect() as db:
        batch = db.execute("SELECT * FROM invoice_import_batches WHERE token=? AND workspace=?", (token, workspace)).fetchone()
        if not batch:
            raise ValueError("The invoice preview was not found. Upload the file again.")
        if batch["status"] != "preview":
            raise ValueError("This invoice preview has already been confirmed")
        if int(batch["user_id"]) != int(user["id"]) and not auth.can_admin(user):
            raise PermissionError("This invoice preview belongs to another user")
        records = json.loads(batch["payload_json"])
        grouped: dict[int, list[dict[str, Any]]] = {}
        for record in records:
            grouped.setdefault(int(record["quote_id"]), []).append(record)

        for quote_id, incoming in grouped.items():
            quote = db.execute("SELECT * FROM quotes WHERE id=?", (quote_id,)).fetchone()
            # CAMBIO 3: Quitamos la restricción de po_detected al confirmar
            if not quote or quote["status"] != "pending" or quote["po_date"]:
                continue
            
            current = po_invoices.active(db, "quote_invoices", quote_id)
            actual = [row for row in current if not int(row.get("is_legacy") or 0)]
            if len(actual) != len(current):
                legacy_replaced += len(current) - len(actual)
            keys = {(_normalize(row["invoice_series"]), _normalize(row["invoice_number"])) for row in actual}
            added_for_quote: list[dict[str, Any]] = []
            for row in incoming:
                key = (_normalize(row["invoice_series"]), _normalize(row["invoice_number"]))
                if key in keys:
                    continue
                keys.add(key)
                added_for_quote.append(
                    {
                        "invoice_date": row["invoice_date"],
                        "invoice_series": row["invoice_series"],
                        "invoice_number": row["invoice_number"],
                        "amount": row["amount"],
                    }
                )
            if not added_for_quote:
                continue
            combined = po_invoices.normalize(actual + added_for_quote, "USD")
            po_invoices.replace(db, "quote_invoices", quote_id, combined, user["id"], user["display_name"], timestamp)
            total, po_date = po_invoices.totals(combined)
            db.execute(
                """UPDATE quotes SET status='po',is_safe=0,loss_reason='',po_total_usd=?,po_date=?,
                updated_at=?,status_changed_at=? WHERE id=?""",
                (total, po_date, timestamp, timestamp, quote_id),
            )
            note = po_invoices.automatic_comment(combined, "USD", str(user.get("language") or "en"), updated=False)
            db.execute(
                """INSERT INTO quote_comments(quote_id,body,user_id,user_name,created_at)
                VALUES(?,?,?,?,?)""", (quote_id, note, user["id"], user["display_name"], timestamp),
            )
            db.execute(
                """INSERT INTO quote_events(quote_id,event_type,follow_up_type,note,user_id,user_name,created_at)
                VALUES(?,'po_invoices_registered',?,?,?,?,?)""",
                (quote_id, quote["follow_up_type"] or "", note, user["id"], user["display_name"], timestamp),
            )
            db.execute(
                """INSERT INTO quote_events(quote_id,event_type,from_status,to_status,follow_up_type,note,
                user_id,user_name,created_at) VALUES(?,'status_changed','pending','po',?,'',?,?,?)""",
                (quote_id, quote["follow_up_type"] or "", user["id"], user["display_name"], timestamp),
            )
            invoices_added += len(added_for_quote)
            quotes_updated += 1

            db.execute("INSERT INTO po_invoices (quote_id) SELECT ? WHERE NOT EXIST(SELECT 1 FROM po_invoices WHERE quote_id?)", (quote_id, quote_id))

        db.execute("UPDATE invoice_import_batches SET status='confirmed',confirmed_at=? WHERE token=?", (timestamp, token))

    detail = f"{batch['filename']}: {invoices_added} invoices added to {quotes_updated} quotations"
    auth.audit(user, "invoice_import_confirmed", workspace, "invoice_import", token, detail)
    return {
        "ok": True,
        "invoices_added": invoices_added,
        "quotes_updated": quotes_updated,
        "legacy_invoices_replaced": legacy_replaced,
        "message": f"{invoices_added} invoices were added to {quotes_updated} PO quotations.",
    }